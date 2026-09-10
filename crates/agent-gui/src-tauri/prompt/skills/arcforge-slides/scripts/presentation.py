#!/usr/bin/env python3
"""Deterministic PPTX creation, inspection, and optional PDF rendering for ArcForge.

Two deck specifications are supported:

* the legacy ``slides[]`` layout specification (``schema_version`` 1 or absent), and
* the ``schema_version: 3`` deck manifest whose pages are SVG page descriptions
  converted into native PPTX objects (shapes, text boxes, pictures, charts).
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence, Tuple

try:
    from lxml import etree
    from PIL import Image, ImageFont
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Inches, Pt

    PPTX_IMPORT_ERROR: Optional[Exception] = None
except (
    Exception
) as error:  # pragma: no cover - exercised only when dependency is absent
    PPTX_IMPORT_ERROR = error


DEFAULT_THEME = {
    "background": "F8FAFC",
    "surface": "E2E8F0",
    "primary": "0F172A",
    "accent": "2563EB",
    "text": "0F172A",
    "muted": "64748B",
    "inverse": "F8FAFC",
    "font": "Microsoft YaHei",
    "font_alt": "Aptos",
}
SERIES_COLORS = ["2563EB", "0F172A", "0891B2", "7C3AED", "EA580C", "16A34A"]
SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5


class PresentationError(RuntimeError):
    pass


def require_pptx() -> None:
    if PPTX_IMPORT_ERROR is not None:
        raise PresentationError(
            "python-pptx and Pillow are required. Ask the user before installing "
            "scripts/requirements.txt. Import error: " + str(PPTX_IMPORT_ERROR)
        )


def load_json_object(path_value: str) -> Tuple[Dict[str, Any], Path]:
    path = Path(path_value).expanduser().resolve()
    if not path.is_file():
        raise PresentationError("JSON specification does not exist: " + str(path))
    try:
        with path.open("r", encoding="utf-8-sig") as handle:
            value = json.load(handle)
    except (OSError, json.JSONDecodeError) as error:
        raise PresentationError(
            "Failed to read JSON specification: " + str(error)
        ) from error
    if not isinstance(value, dict):
        raise PresentationError("JSON specification must be an object")
    return value, path


def normalize_color(value: Any, field: str = "color") -> str:
    raw = str(value).strip().lstrip("#").upper()
    if len(raw) != 6 or any(char not in "0123456789ABCDEF" for char in raw):
        raise PresentationError(field + " must be a 6-digit hexadecimal color")
    return raw


def merged_theme(spec: Mapping[str, Any]) -> Dict[str, str]:
    theme = dict(DEFAULT_THEME)
    raw = spec.get("theme")
    if raw is None:
        return theme
    if not isinstance(raw, dict):
        raise PresentationError("theme must be an object")
    for key in (
        "background",
        "surface",
        "primary",
        "accent",
        "text",
        "muted",
        "inverse",
    ):
        if key in raw:
            theme[key] = normalize_color(raw[key], "theme." + key)
    for key in ("font", "font_alt"):
        if key in raw and str(raw[key]).strip():
            theme[key] = str(raw[key]).strip()
    return theme


def rgb(value: str) -> Any:
    return RGBColor.from_string(normalize_color(value))


def set_slide_background(slide: Any, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_solid_shape(
    slide: Any,
    shape_type: Any,
    x: float,
    y: float,
    width: float,
    height: float,
    fill_color: str,
    line_color: Optional[str] = None,
    radius_name: Optional[str] = None,
) -> Any:
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_color)
    if line_color:
        shape.line.color.rgb = rgb(line_color)
    else:
        shape.line.fill.background()
    if radius_name:
        shape.name = radius_name
    return shape


def add_text(
    slide: Any,
    text: Any,
    x: float,
    y: float,
    width: float,
    height: float,
    theme: Mapping[str, str],
    font_size: float = 18,
    color: Optional[str] = None,
    bold: bool = False,
    align: Any = None,
    vertical: Any = None,
    font_name: Optional[str] = None,
    margin: float = 0.02,
    name: Optional[str] = None,
) -> Any:
    shape = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    if vertical is not None:
        frame.vertical_anchor = vertical
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align if align is not None else PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = font_name or theme["font"]
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color or theme["text"])
    return shape


def add_content_header(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> float:
    set_slide_background(slide, str(slide_spec.get("background", theme["background"])))
    accent = normalize_color(slide_spec.get("accent", theme["accent"]), "slide.accent")
    add_solid_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, 0.09, accent)
    title = str(slide_spec.get("title", "")).strip()
    if not title:
        raise PresentationError("Content slide requires a title")
    add_text(
        slide,
        title,
        0.7,
        0.38,
        11.9,
        0.62,
        theme,
        font_size=26,
        bold=True,
        name="ArcForge Title",
    )
    subtitle = str(slide_spec.get("subtitle", "")).strip()
    if subtitle:
        add_text(
            slide,
            subtitle,
            0.72,
            1.0,
            11.6,
            0.35,
            theme,
            font_size=11.5,
            color=theme["muted"],
        )
        return 1.52
    return 1.30


def add_footer(
    slide: Any,
    footer_text: str,
    slide_number: int,
    theme: Mapping[str, str],
) -> None:
    add_solid_shape(
        slide,
        MSO_SHAPE.RECTANGLE,
        0.7,
        6.94,
        11.93,
        0.012,
        theme["surface"],
    )
    if footer_text:
        add_text(
            slide,
            footer_text,
            0.7,
            7.01,
            10.9,
            0.22,
            theme,
            font_size=8.5,
            color=theme["muted"],
        )
    add_text(
        slide,
        str(slide_number),
        11.75,
        7.01,
        0.85,
        0.22,
        theme,
        font_size=8.5,
        color=theme["muted"],
        align=PP_ALIGN.RIGHT,
    )


def normalized_bullets(raw: Any) -> List[Dict[str, Any]]:
    if not isinstance(raw, list):
        raise PresentationError("bullets must be an array")
    bullets: List[Dict[str, Any]] = []
    for item in raw:
        if isinstance(item, str):
            bullets.append({"text": item, "level": 0})
        elif isinstance(item, dict) and str(item.get("text", "")).strip():
            bullets.append(
                {
                    "text": str(item["text"]),
                    "level": max(0, min(3, int(item.get("level", 0)))),
                    "accent": item.get("accent"),
                }
            )
        else:
            raise PresentationError("each bullet must be text or an object with text")
    return bullets


def add_bullet_list(
    slide: Any,
    bullets: Sequence[Mapping[str, Any]],
    x: float,
    y: float,
    width: float,
    height: float,
    theme: Mapping[str, str],
    font_size: float = 18,
) -> None:
    if not bullets:
        return
    row_height = min(0.72, height / max(1, len(bullets)))
    for index, item in enumerate(bullets):
        level = int(item.get("level", 0))
        item_y = y + index * row_height
        indent = level * 0.30
        marker_color = (
            normalize_color(item["accent"], "bullet.accent")
            if item.get("accent")
            else theme["accent"]
        )
        add_solid_shape(
            slide,
            MSO_SHAPE.OVAL,
            x + indent,
            item_y + 0.18,
            0.10,
            0.10,
            marker_color,
        )
        add_text(
            slide,
            item["text"],
            x + 0.24 + indent,
            item_y,
            width - 0.24 - indent,
            row_height,
            theme,
            font_size=max(12, font_size - level * 1.5),
            vertical=MSO_ANCHOR.MIDDLE,
        )


def render_title_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    set_slide_background(slide, str(slide_spec.get("background", theme["primary"])))
    accent = normalize_color(slide_spec.get("accent", theme["accent"]), "slide.accent")
    add_solid_shape(slide, MSO_SHAPE.RECTANGLE, 0.78, 1.22, 0.12, 4.75, accent)
    kicker = str(slide_spec.get("kicker", "")).strip()
    if kicker:
        add_text(
            slide,
            kicker.upper(),
            1.18,
            1.18,
            10.9,
            0.38,
            theme,
            font_size=10,
            color=accent,
            bold=True,
        )
    title = str(slide_spec.get("title", "")).strip()
    if not title:
        raise PresentationError("title slide requires a title")
    add_text(
        slide,
        title,
        1.16,
        1.72,
        10.9,
        2.0,
        theme,
        font_size=34,
        color=theme["inverse"],
        bold=True,
        vertical=MSO_ANCHOR.MIDDLE,
        name="ArcForge Title",
    )
    subtitle = str(slide_spec.get("subtitle", "")).strip()
    if subtitle:
        add_text(
            slide,
            subtitle,
            1.18,
            4.10,
            9.9,
            1.0,
            theme,
            font_size=17,
            color="CBD5E1",
        )
    footer = str(slide_spec.get("footer", "ArcForge")).strip()
    add_text(
        slide,
        footer,
        1.18,
        6.67,
        10.0,
        0.32,
        theme,
        font_size=9,
        color="94A3B8",
    )


def render_section_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    set_slide_background(slide, str(slide_spec.get("background", theme["primary"])))
    accent = normalize_color(slide_spec.get("accent", theme["accent"]), "slide.accent")
    number = str(slide_spec.get("number", "")).strip()
    if number:
        add_text(
            slide,
            number,
            0.85,
            0.62,
            2.3,
            1.3,
            theme,
            font_size=58,
            color=accent,
            bold=True,
        )
    title = str(slide_spec.get("title", "")).strip()
    if not title:
        raise PresentationError("section slide requires a title")
    add_text(
        slide,
        title,
        0.88,
        2.18,
        11.2,
        1.45,
        theme,
        font_size=34,
        color=theme["inverse"],
        bold=True,
        vertical=MSO_ANCHOR.MIDDLE,
        name="ArcForge Title",
    )
    subtitle = str(slide_spec.get("subtitle", "")).strip()
    if subtitle:
        add_text(
            slide,
            subtitle,
            0.90,
            3.92,
            10.8,
            0.75,
            theme,
            font_size=16,
            color="CBD5E1",
        )
    add_solid_shape(slide, MSO_SHAPE.RECTANGLE, 0.88, 5.22, 2.2, 0.08, accent)


def render_bullets_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    content_top = add_content_header(slide, slide_spec, theme)
    bullets = normalized_bullets(slide_spec.get("bullets", []))
    if len(bullets) > 8:
        raise PresentationError(
            "bullets slide supports at most 8 bullets; split dense content"
        )
    add_bullet_list(slide, bullets, 0.92, content_top + 0.12, 11.25, 5.0, theme, 18)


def render_two_column_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    content_top = add_content_header(slide, slide_spec, theme)
    for index, key in enumerate(("left", "right")):
        column = slide_spec.get(key, {})
        if not isinstance(column, dict):
            raise PresentationError(key + " must be an object")
        x = 0.70 if index == 0 else 6.72
        add_solid_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            content_top + 0.08,
            5.90,
            4.92,
            "FFFFFF",
            line_color=theme["surface"],
        )
        heading = str(column.get("heading", key.title())).strip()
        add_text(
            slide,
            heading,
            x + 0.38,
            content_top + 0.35,
            5.10,
            0.48,
            theme,
            font_size=17,
            bold=True,
        )
        bullets = normalized_bullets(column.get("bullets", []))
        if len(bullets) > 6:
            raise PresentationError(key + " column supports at most 6 bullets")
        add_bullet_list(
            slide,
            bullets,
            x + 0.40,
            content_top + 1.02,
            5.00,
            3.55,
            theme,
            14,
        )


def render_metrics_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    content_top = add_content_header(slide, slide_spec, theme)
    metrics = slide_spec.get("metrics")
    if not isinstance(metrics, list) or not metrics:
        raise PresentationError("metrics slide requires a non-empty metrics array")
    if len(metrics) > 4:
        raise PresentationError("metrics slide supports at most 4 metric cards")
    gap = 0.24
    total_width = 11.93
    card_width = (total_width - gap * (len(metrics) - 1)) / len(metrics)
    for index, metric in enumerate(metrics):
        if not isinstance(metric, dict):
            raise PresentationError("each metric must be an object")
        x = 0.70 + index * (card_width + gap)
        add_solid_shape(
            slide,
            MSO_SHAPE.ROUNDED_RECTANGLE,
            x,
            content_top + 0.35,
            card_width,
            3.72,
            "FFFFFF",
            line_color=theme["surface"],
        )
        add_solid_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            x,
            content_top + 0.35,
            card_width,
            0.09,
            theme["accent"],
        )
        add_text(
            slide,
            metric.get("value", ""),
            x + 0.28,
            content_top + 0.90,
            card_width - 0.56,
            0.90,
            theme,
            font_size=28,
            bold=True,
            vertical=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            metric.get("label", ""),
            x + 0.28,
            content_top + 1.92,
            card_width - 0.56,
            0.90,
            theme,
            font_size=12,
            color=theme["muted"],
        )
        if metric.get("delta") is not None:
            add_text(
                slide,
                metric["delta"],
                x + 0.28,
                content_top + 3.08,
                card_width - 0.56,
                0.34,
                theme,
                font_size=11,
                color=theme["accent"],
                bold=True,
            )


def set_table_cell(
    cell: Any,
    value: Any,
    theme: Mapping[str, str],
    fill_color: str,
    text_color: str,
    bold: bool,
    font_size: float,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill_color)
    cell.margin_left = Inches(0.08)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.05)
    cell.margin_bottom = Inches(0.05)
    frame = cell.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = str(value)
    run.font.name = theme["font"]
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = rgb(text_color)


def render_table_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    content_top = add_content_header(slide, slide_spec, theme)
    columns = slide_spec.get("columns")
    rows = slide_spec.get("rows")
    if not isinstance(columns, list) or not columns:
        raise PresentationError("table slide requires a non-empty columns array")
    if not isinstance(rows, list):
        raise PresentationError("table slide rows must be an array")
    if len(columns) > 8 or len(rows) > 12:
        raise PresentationError(
            "table exceeds the supported 8 columns or 12 rows; split it"
        )
    total_rows = len(rows) + 1
    x, y, width, height = 0.70, content_top + 0.18, 11.93, 5.12
    table_shape = slide.shapes.add_table(
        total_rows,
        len(columns),
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(height),
    )
    table_shape.name = "ArcForge Table"
    table = table_shape.table

    raw_widths = slide_spec.get("column_widths")
    if raw_widths is not None:
        if not isinstance(raw_widths, list) or len(raw_widths) != len(columns):
            raise PresentationError("column_widths must match the number of columns")
        weights = [float(item) for item in raw_widths]
        if any(item <= 0 for item in weights):
            raise PresentationError("column_widths values must be positive")
    else:
        weights = [1.0] * len(columns)
    weight_total = sum(weights)
    for index, weight in enumerate(weights):
        table.columns[index].width = Inches(width * weight / weight_total)

    for column_index, value in enumerate(columns):
        set_table_cell(
            table.cell(0, column_index),
            value,
            theme,
            theme["primary"],
            theme["inverse"],
            True,
            11,
        )
    body_font_size = 10.5 if len(rows) <= 8 else 9
    for row_index, raw_row in enumerate(rows, start=1):
        if not isinstance(raw_row, list):
            raise PresentationError("each table row must be an array")
        values = list(raw_row[: len(columns)]) + [""] * max(
            0, len(columns) - len(raw_row)
        )
        fill_color = "FFFFFF" if row_index % 2 else "F1F5F9"
        for column_index, value in enumerate(values):
            set_table_cell(
                table.cell(row_index, column_index),
                value,
                theme,
                fill_color,
                theme["text"],
                False,
                body_font_size,
            )


def chart_type_value(value: str) -> Any:
    normalized = value.lower()
    mapping = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
    }
    if normalized not in mapping:
        raise PresentationError("Unsupported chart_type: " + value)
    return mapping[normalized]


def render_chart_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    content_top = add_content_header(slide, slide_spec, theme)
    categories = slide_spec.get("categories")
    series = slide_spec.get("series")
    if not isinstance(categories, list) or not categories:
        raise PresentationError("chart slide requires a non-empty categories array")
    if not isinstance(series, list) or not series:
        raise PresentationError("chart slide requires a non-empty series array")

    chart_data = ChartData()
    chart_data.categories = [str(item) for item in categories]
    for item in series:
        if not isinstance(item, dict) or not str(item.get("name", "")).strip():
            raise PresentationError("each chart series requires a name and values")
        values = item.get("values")
        if not isinstance(values, list) or len(values) != len(categories):
            raise PresentationError(
                "each chart series values array must match categories"
            )
        chart_data.add_series(str(item["name"]), [float(value) for value in values])

    chart_type = str(slide_spec.get("chart_type", "column"))
    chart_frame = slide.shapes.add_chart(
        chart_type_value(chart_type),
        Inches(0.82),
        Inches(content_top + 0.12),
        Inches(11.68),
        Inches(5.08),
        chart_data,
    )
    chart_frame.name = "ArcForge Chart"
    chart = chart_frame.chart
    chart.has_title = False
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    chart.chart_style = int(slide_spec.get("chart_style", 10))

    for index, chart_series in enumerate(chart.series):
        color = rgb(SERIES_COLORS[index % len(SERIES_COLORS)])
        if chart_type.lower() == "line":
            chart_series.format.line.color.rgb = color
            chart_series.format.line.width = Pt(2.25)
        else:
            chart_series.format.fill.solid()
            chart_series.format.fill.fore_color.rgb = color
            chart_series.format.line.color.rgb = color

    if chart_type.lower() not in ("pie", "doughnut"):
        chart.category_axis.tick_labels.font.name = theme["font_alt"]
        chart.category_axis.tick_labels.font.size = Pt(10)
        chart.value_axis.tick_labels.font.name = theme["font_alt"]
        chart.value_axis.tick_labels.font.size = Pt(10)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = rgb(theme["surface"])
    else:
        plot = chart.plots[0]
        plot.has_data_labels = True
        plot.data_labels.show_percentage = True
        plot.data_labels.show_legend_key = False


def resolve_asset_path(raw_path: Any, spec_dir: Path) -> Path:
    path = Path(str(raw_path)).expanduser()
    if not path.is_absolute():
        path = spec_dir / path
    path = path.resolve()
    if not path.is_file():
        raise PresentationError("Image does not exist: " + str(path))
    return path


def add_picture_fit(
    slide: Any,
    image_path: Path,
    x: float,
    y: float,
    width: float,
    height: float,
    fit: str,
) -> Any:
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    if image_width <= 0 or image_height <= 0:
        raise PresentationError("Image has invalid dimensions: " + str(image_path))
    image_ratio = image_width / image_height
    frame_ratio = width / height

    if fit == "contain":
        if image_ratio >= frame_ratio:
            fitted_width = width
            fitted_height = width / image_ratio
        else:
            fitted_height = height
            fitted_width = height * image_ratio
        left = x + (width - fitted_width) / 2
        top = y + (height - fitted_height) / 2
        picture = slide.shapes.add_picture(
            str(image_path),
            Inches(left),
            Inches(top),
            Inches(fitted_width),
            Inches(fitted_height),
        )
    elif fit == "cover":
        picture = slide.shapes.add_picture(
            str(image_path),
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height),
        )
        if image_ratio > frame_ratio:
            crop = (1.0 - frame_ratio / image_ratio) / 2.0
            picture.crop_left = crop
            picture.crop_right = crop
        elif image_ratio < frame_ratio:
            crop = (1.0 - image_ratio / frame_ratio) / 2.0
            picture.crop_top = crop
            picture.crop_bottom = crop
    else:
        raise PresentationError("image fit must be contain or cover")
    picture.name = "ArcForge Image"
    return picture


def render_image_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
    spec_dir: Path,
) -> None:
    content_top = add_content_header(slide, slide_spec, theme)
    if not slide_spec.get("image"):
        raise PresentationError("image slide requires an image path")
    image_path = resolve_asset_path(slide_spec["image"], spec_dir)
    caption = str(slide_spec.get("caption", "")).strip()
    box_height = 4.55 if caption else 5.12
    add_solid_shape(
        slide,
        MSO_SHAPE.ROUNDED_RECTANGLE,
        0.70,
        content_top + 0.10,
        11.93,
        box_height,
        "FFFFFF",
        line_color=theme["surface"],
    )
    add_picture_fit(
        slide,
        image_path,
        0.84,
        content_top + 0.24,
        11.65,
        box_height - 0.28,
        str(slide_spec.get("fit", "contain")).lower(),
    )
    if caption:
        add_text(
            slide,
            caption,
            0.82,
            content_top + 4.83,
            11.55,
            0.34,
            theme,
            font_size=9.5,
            color=theme["muted"],
            align=PP_ALIGN.CENTER,
        )


def render_quote_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    set_slide_background(slide, str(slide_spec.get("background", theme["background"])))
    quote = str(slide_spec.get("quote", "")).strip()
    if not quote:
        raise PresentationError("quote slide requires quote text")
    title = str(slide_spec.get("title", "")).strip()
    if title:
        add_text(
            slide,
            title,
            0.72,
            0.45,
            11.7,
            0.55,
            theme,
            font_size=18,
            color=theme["muted"],
            name="ArcForge Title",
        )
    add_text(
        slide,
        "“",
        0.82,
        1.30,
        1.0,
        1.0,
        theme,
        font_size=68,
        color=theme["accent"],
        bold=True,
    )
    add_text(
        slide,
        quote,
        1.48,
        1.72,
        10.35,
        3.10,
        theme,
        font_size=27,
        bold=True,
        vertical=MSO_ANCHOR.MIDDLE,
        name=None if title else "ArcForge Title",
    )
    attribution = str(slide_spec.get("attribution", "")).strip()
    if attribution:
        add_text(
            slide,
            "— " + attribution,
            1.52,
            5.12,
            9.6,
            0.55,
            theme,
            font_size=13,
            color=theme["muted"],
        )


def render_closing_slide(
    slide: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
) -> None:
    set_slide_background(slide, str(slide_spec.get("background", theme["primary"])))
    accent = normalize_color(slide_spec.get("accent", theme["accent"]), "slide.accent")
    add_solid_shape(slide, MSO_SHAPE.RECTANGLE, 0.82, 1.25, 0.10, 4.60, accent)
    title = str(slide_spec.get("title", "")).strip()
    if not title:
        raise PresentationError("closing slide requires a title")
    add_text(
        slide,
        title,
        1.22,
        1.65,
        10.6,
        1.65,
        theme,
        font_size=31,
        color=theme["inverse"],
        bold=True,
        vertical=MSO_ANCHOR.MIDDLE,
        name="ArcForge Title",
    )
    subtitle = str(slide_spec.get("subtitle", "")).strip()
    if subtitle:
        add_text(
            slide,
            subtitle,
            1.24,
            3.72,
            9.9,
            0.92,
            theme,
            font_size=16,
            color="CBD5E1",
        )
    contact = str(slide_spec.get("contact", "")).strip()
    if contact:
        add_text(
            slide,
            contact,
            1.24,
            5.42,
            9.8,
            0.42,
            theme,
            font_size=11,
            color=accent,
            bold=True,
        )


def apply_notes(slide: Any, notes: Any) -> None:
    if notes is None or not str(notes).strip():
        return
    try:
        slide.notes_slide.notes_text_frame.text = str(notes)
    except Exception as error:
        raise PresentationError("Failed to write slide notes: " + str(error)) from error


def render_slide(
    presentation: Any,
    slide_spec: Mapping[str, Any],
    theme: Mapping[str, str],
    spec_dir: Path,
    footer_text: str,
    slide_number: int,
) -> None:
    if not isinstance(slide_spec, dict):
        raise PresentationError("each slides entry must be an object")
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_type = str(slide_spec.get("type", "bullets")).strip().lower()
    if slide_type == "title":
        render_title_slide(slide, slide_spec, theme)
    elif slide_type == "section":
        render_section_slide(slide, slide_spec, theme)
    elif slide_type == "bullets":
        render_bullets_slide(slide, slide_spec, theme)
    elif slide_type == "two-column":
        render_two_column_slide(slide, slide_spec, theme)
    elif slide_type == "metrics":
        render_metrics_slide(slide, slide_spec, theme)
    elif slide_type == "table":
        render_table_slide(slide, slide_spec, theme)
    elif slide_type == "chart":
        render_chart_slide(slide, slide_spec, theme)
    elif slide_type == "image":
        render_image_slide(slide, slide_spec, theme, spec_dir)
    elif slide_type == "quote":
        render_quote_slide(slide, slide_spec, theme)
    elif slide_type == "closing":
        render_closing_slide(slide, slide_spec, theme)
    else:
        raise PresentationError("Unsupported slide type: " + slide_type)
    if slide_type not in ("title", "section", "closing"):
        add_footer(slide, footer_text, slide_number, theme)
    apply_notes(slide, slide_spec.get("notes"))


def apply_metadata(presentation: Any, spec: Mapping[str, Any]) -> None:
    metadata = spec.get("metadata", {})
    if not isinstance(metadata, dict):
        raise PresentationError("metadata must be an object")
    properties = presentation.core_properties
    mapping = {
        "title": "title",
        "subject": "subject",
        "author": "author",
        "comments": "comments",
        "keywords": "keywords",
        "category": "category",
    }
    for source, target in mapping.items():
        if source in metadata:
            setattr(properties, target, str(metadata[source]))
    if not properties.author:
        properties.author = "ArcForge"
    if metadata.get("company") and not properties.comments:
        properties.comments = "Company: " + str(metadata["company"])


def create_presentation(spec: Mapping[str, Any], spec_dir: Path) -> Any:
    slides = spec.get("slides")
    if not isinstance(slides, list) or not slides:
        raise PresentationError("slides must be a non-empty array")
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH)
    presentation.slide_height = Inches(SLIDE_HEIGHT)
    apply_metadata(presentation, spec)
    theme = merged_theme(spec)
    footer = str(spec.get("footer", "")).strip()
    for index, slide_spec in enumerate(slides, start=1):
        render_slide(presentation, slide_spec, theme, spec_dir, footer, index)
    return presentation


# ---------------------------------------------------------------------------
# schema_version 3: SVG page descriptions converted into native PPTX objects
# ---------------------------------------------------------------------------

SVG_NS = "{http://www.w3.org/2000/svg}"
XLINK_NS = "{http://www.w3.org/1999/xlink}"
CANVAS_WIDTH = 1280.0
CANVAS_HEIGHT = 720.0
EMU_PER_PX = 9525
MAX_SVG_BYTES = 2 * 1024 * 1024
MAX_SVG_ELEMENTS = 2000
DECK_SCHEMA_VERSION = 3
SUPPORTED_SVG_ELEMENTS = frozenset(
    {
        "svg",
        "g",
        "defs",
        "linearGradient",
        "stop",
        "rect",
        "circle",
        "ellipse",
        "line",
        "polygon",
        "path",
        "text",
        "tspan",
        "image",
        "title",
        "desc",
    }
)
SUPPORTED_SVG_ATTRIBUTES = frozenset(
    {
        "id",
        "x",
        "y",
        "x1",
        "y1",
        "x2",
        "y2",
        "cx",
        "cy",
        "r",
        "rx",
        "ry",
        "dx",
        "dy",
        "width",
        "height",
        "d",
        "points",
        "fill",
        "fill-opacity",
        "stroke",
        "stroke-width",
        "font-size",
        "font-weight",
        "font-family",
        "text-anchor",
        "transform",
        "viewBox",
        "xmlns",
        "preserveAspectRatio",
        "offset",
        "stop-color",
        "stop-opacity",
        "gradientUnits",
        "data-arcforge",
        "data-asset",
        "data-chart",
        "data-width",
        "data-fill",
        "data-stroke",
        "data-render",
        "data-role",
        "role",
        "lang",
    }
)
CHART_TYPES_V3 = {
    "column": "column",
    "bar": "bar",
    "line": "line",
    "pie": "pie",
    "doughnut": "doughnut",
    "area": "area",
}
FALLBACK_FONTS = ("Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "PingFang SC", "Arial")


def px_to_emu(value: float) -> Any:
    return Emu(int(round(value * EMU_PER_PX)))


def svg_local_name(tag: str) -> str:
    return tag.split("}")[-1]


def parse_svg_length(value: Any, label: str) -> float:
    text = str(value).strip()
    if text.endswith("px"):
        text = text[:-2]
    try:
        return float(text)
    except ValueError as error:
        raise PresentationError(label + " must be a number in px, got " + str(value)) from error


def parse_svg_color(value: Optional[str], label: str) -> Optional[str]:
    if value is None:
        return None
    text = value.strip()
    if text in ("none", "transparent", ""):
        return None
    if text.startswith("#"):
        digits = text[1:]
        if len(digits) == 3:
            digits = "".join(character * 2 for character in digits)
        if len(digits) == 8:
            digits = digits[:6]
        return normalize_color(digits, label)
    match = re.match(r"rgb\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*\)", text)
    if match:
        return "".join("%02X" % min(255, int(part)) for part in match.groups())
    named = {
        "white": "FFFFFF",
        "black": "000000",
        "gray": "808080",
        "grey": "808080",
        "red": "FF0000",
        "blue": "0000FF",
    }
    if text.lower() in named:
        return named[text.lower()]
    raise PresentationError(label + " uses unsupported color " + text)


@dataclass
class SvgStyle:
    fill: Optional[str] = "000000"
    fill_opacity: float = 1.0
    stroke: Optional[str] = None
    stroke_width: float = 1.0
    font_family: str = "Microsoft YaHei"
    font_size: float = 16.0
    font_weight: str = "normal"
    text_anchor: str = "start"
    dx: float = 0.0
    dy: float = 0.0

    def inherit(self, element: Any, page: str) -> "SvgStyle":
        style = SvgStyle(**self.__dict__)
        attributes = dict(element.attrib)
        if "style" in attributes:
            raise PresentationError(
                page + ": inline style attributes are not supported; use presentation attributes"
            )
        for key in attributes:
            if key.startswith("{"):
                if key.startswith(XLINK_NS):
                    raise PresentationError(page + ": xlink attributes are not supported")
                continue
            if key not in SUPPORTED_SVG_ATTRIBUTES:
                raise PresentationError(
                    page + ": unsupported SVG attribute '" + key + "' on <" + svg_local_name(element.tag) + ">"
                )
        if "fill" in attributes:
            raw = attributes["fill"].strip()
            style.fill = raw if raw.startswith("url(") else parse_svg_color(raw, page + " fill")
        if "fill-opacity" in attributes:
            style.fill_opacity = max(0.0, min(1.0, parse_svg_length(attributes["fill-opacity"], page + " fill-opacity")))
        if "stroke" in attributes:
            style.stroke = parse_svg_color(attributes["stroke"], page + " stroke")
        if "stroke-width" in attributes:
            style.stroke_width = parse_svg_length(attributes["stroke-width"], page + " stroke-width")
        if "font-family" in attributes:
            style.font_family = attributes["font-family"].split(",")[0].strip("'\" ") or style.font_family
        if "font-size" in attributes:
            style.font_size = parse_svg_length(attributes["font-size"], page + " font-size")
        if "font-weight" in attributes:
            style.font_weight = attributes["font-weight"].strip().lower()
        if "text-anchor" in attributes:
            style.text_anchor = attributes["text-anchor"].strip().lower()
        if "transform" in attributes:
            match = re.fullmatch(
                r"\s*translate\(\s*(-?[\d.]+)(?:[\s,]+(-?[\d.]+))?\s*\)\s*",
                attributes["transform"],
            )
            if not match:
                raise PresentationError(
                    page + ": only transform=\"translate(x y)\" is supported, got " + attributes["transform"]
                )
            style.dx += float(match.group(1))
            style.dy += float(match.group(2) or 0.0)
        return style


@dataclass
class DeckFonts:
    heading: List[str] = field(default_factory=lambda: list(FALLBACK_FONTS))
    body: List[str] = field(default_factory=lambda: list(FALLBACK_FONTS))


def font_directories() -> List[Path]:
    directories: List[Path] = []
    windir = os.environ.get("WINDIR") or os.environ.get("SystemRoot")
    if windir:
        directories.append(Path(windir) / "Fonts")
    local_app = os.environ.get("LOCALAPPDATA")
    if local_app:
        directories.append(Path(local_app) / "Microsoft" / "Windows" / "Fonts")
    directories.extend(
        [
            Path("/usr/share/fonts"),
            Path("/usr/local/share/fonts"),
            Path.home() / ".fonts",
            Path("/System/Library/Fonts"),
            Path("/Library/Fonts"),
        ]
    )
    return [directory for directory in directories if directory.is_dir()]


FONT_FILE_CANDIDATES: Dict[Tuple[str, bool], Tuple[str, ...]] = {
    ("microsoft yahei", False): ("msyh.ttc", "msyh.ttf"),
    ("microsoft yahei", True): ("msyhbd.ttc", "msyhbd.ttf", "msyh.ttc"),
    ("simhei", False): ("simhei.ttf",),
    ("simhei", True): ("simhei.ttf",),
    ("simsun", False): ("simsun.ttc",),
    ("simsun", True): ("simsun.ttc",),
    ("dengxian", False): ("Deng.ttf",),
    ("dengxian", True): ("Dengb.ttf", "Deng.ttf"),
    ("arial", False): ("arial.ttf",),
    ("arial", True): ("arialbd.ttf", "arial.ttf"),
    ("noto sans cjk sc", False): ("NotoSansCJK-Regular.ttc", "NotoSansCJKsc-Regular.otf"),
    ("noto sans cjk sc", True): ("NotoSansCJK-Bold.ttc", "NotoSansCJKsc-Bold.otf"),
}


class TextMeasurer:
    """Measure rendered text width with real font files when available."""

    def __init__(self) -> None:
        self._cache: Dict[Tuple[str, bool, int], Any] = {}
        self._resolved: Dict[Tuple[str, bool], Optional[Path]] = {}
        self.missing_fonts: List[str] = []
        self.measured_with_fonts = True

    def _resolve_file(self, family: str, bold: bool) -> Optional[Path]:
        key = (family.lower(), bold)
        if key in self._resolved:
            return self._resolved[key]
        candidates = FONT_FILE_CANDIDATES.get(key) or FONT_FILE_CANDIDATES.get((key[0], False)) or ()
        found: Optional[Path] = None
        for directory in font_directories():
            for name in candidates:
                candidate = directory / name
                if candidate.is_file():
                    found = candidate
                    break
            if found:
                break
        self._resolved[key] = found
        return found

    def _font(self, family: str, bold: bool, size: float) -> Optional[Any]:
        size_px = max(1, int(round(size)))
        key = (family.lower(), bold, size_px)
        if key in self._cache:
            return self._cache[key]
        font = None
        for candidate_family in [family] + [name for name in FALLBACK_FONTS if name.lower() != family.lower()]:
            path = self._resolve_file(candidate_family, bold)
            if path is None:
                continue
            try:
                font = ImageFont.truetype(str(path), size_px)
            except OSError:
                continue
            if candidate_family.lower() != family.lower() and family not in self.missing_fonts:
                self.missing_fonts.append(family)
            break
        if font is None:
            self.measured_with_fonts = False
            if family not in self.missing_fonts:
                self.missing_fonts.append(family)
        self._cache[key] = font
        return font

    def width(self, text: str, family: str, bold: bool, size: float) -> float:
        font = self._font(family, bold, size)
        if font is None:
            cjk = sum(1 for character in text if ord(character) > 0x2E7F)
            return cjk * size + (len(text) - cjk) * size * 0.55
        return float(font.getlength(text))

    def wrap(self, text: str, family: str, bold: bool, size: float, width: float) -> List[str]:
        lines: List[str] = []
        current = ""
        for token in re.findall(r"\s+|[A-Za-z0-9_@#%&$€£¥.,:;!?'\"()\[\]/-]+|.", text):
            candidate = current + token
            if not current or self.width(candidate.rstrip(), family, bold, size) <= width:
                current = candidate
                continue
            lines.append(current.rstrip())
            current = token.lstrip()
        if current.strip():
            lines.append(current.rstrip())
        return lines or [""]


@dataclass
class ConversionReport:
    warnings: List[str] = field(default_factory=list)
    text_overflows: List[Dict[str, Any]] = field(default_factory=list)
    out_of_bounds: List[Dict[str, Any]] = field(default_factory=list)
    protected_collisions: List[Dict[str, Any]] = field(default_factory=list)
    missing_fonts: List[str] = field(default_factory=list)
    measured_with_fonts: bool = True
    element_counts: Dict[str, int] = field(default_factory=dict)
    asset_renders: List[Dict[str, Any]] = field(default_factory=list)


@dataclass
class DeckManifest:
    schema_version: int
    mode: str
    stage: str
    template: Optional[Path]
    assets: Dict[str, Path]
    style: Dict[str, Any]
    slides: List[Dict[str, Any]]
    metadata: Dict[str, Any]
    fonts: DeckFonts
    asset_cache: Optional[Path] = None


VECTOR_ASSET_SCHEMA = 2
PLAN_STAGE_GRAY = "8C8C8C"


def load_deck_manifest(
    spec: Mapping[str, Any], spec_dir: Path, asset_cache: Optional[Path] = None
) -> DeckManifest:
    if int(spec.get("schema_version", 1)) != DECK_SCHEMA_VERSION:
        raise PresentationError("schema_version must be " + str(DECK_SCHEMA_VERSION) + " for an SVG deck manifest")
    mode = str(spec.get("mode", "blank")).strip().lower()
    if mode not in ("blank", "template"):
        raise PresentationError("mode must be blank or template")
    stage = str(spec.get("stage", "design")).strip().lower()
    if stage not in ("plan", "design"):
        raise PresentationError("stage must be plan or design")
    template: Optional[Path] = None
    if mode == "template":
        raw_template = spec.get("template")
        if not raw_template:
            raise PresentationError("mode=template requires a template path")
        template = Path(str(raw_template)).expanduser()
        if not template.is_absolute():
            template = spec_dir / template
        template = template.resolve()
        if not template.is_file() or template.suffix.lower() != ".pptx":
            raise PresentationError("template must be an existing .pptx file: " + str(template))
    elif spec.get("template"):
        raise PresentationError("template is only valid when mode is template")

    assets_raw = spec.get("assets", {})
    if assets_raw is None:
        assets_raw = {}
    if not isinstance(assets_raw, dict):
        raise PresentationError("assets must be an object mapping asset ids to workspace image paths")
    assets: Dict[str, Path] = {}
    for asset_id, raw_path in assets_raw.items():
        if not re.fullmatch(r"[A-Za-z0-9_.-]{1,64}", str(asset_id)):
            raise PresentationError("asset id '" + str(asset_id) + "' must use letters, digits, '_', '-' or '.'")
        assets[str(asset_id)] = resolve_asset_path(raw_path, spec_dir)
        if assets[str(asset_id)].suffix.lower() == ".svg" and asset_cache is None:
            raise PresentationError(
                "asset '" + str(asset_id) + "' is an SVG; SVG assets are normalized by ArcForge before conversion, "
                "run create or validate through the OfficeRuntime tool"
            )

    slides_raw = spec.get("slides")
    if not isinstance(slides_raw, list) or not slides_raw:
        raise PresentationError("slides must be a non-empty array of {slide_id, svg} entries")
    seen_ids: set[str] = set()
    slides: List[Dict[str, Any]] = []
    for index, entry in enumerate(slides_raw):
        if not isinstance(entry, dict):
            raise PresentationError("slides[" + str(index) + "] must be an object")
        slide_id = str(entry.get("slide_id", "")).strip()
        if not re.fullmatch(r"[A-Za-z0-9_.-]{1,64}", slide_id):
            raise PresentationError("slides[" + str(index) + "].slide_id must use letters, digits, '_', '-' or '.'")
        if slide_id in seen_ids:
            raise PresentationError("duplicate slide_id " + slide_id)
        seen_ids.add(slide_id)
        raw_svg = entry.get("svg")
        if not raw_svg:
            raise PresentationError("slides[" + str(index) + "].svg is required")
        svg_path = Path(str(raw_svg)).expanduser()
        if not svg_path.is_absolute():
            svg_path = spec_dir / svg_path
        svg_path = svg_path.resolve()
        if not svg_path.is_file() or svg_path.suffix.lower() != ".svg":
            raise PresentationError("slides[" + str(index) + "].svg must be an existing .svg file: " + str(svg_path))
        slides.append(
            {
                "slide_id": slide_id,
                "svg": svg_path,
                "notes": entry.get("notes"),
                "layout": entry.get("layout"),
            }
        )

    style = spec.get("style", {})
    if style is None:
        style = {}
    if not isinstance(style, dict):
        raise PresentationError("style must be an object")
    fonts = DeckFonts()
    raw_fonts = style.get("fonts", {})
    if isinstance(raw_fonts, dict):
        for key in ("heading", "body"):
            value = raw_fonts.get(key)
            if isinstance(value, str) and value.strip():
                setattr(fonts, key, [value.strip()] + [name for name in FALLBACK_FONTS if name != value.strip()])
            elif isinstance(value, list) and value:
                names = [str(item).strip() for item in value if str(item).strip()]
                setattr(fonts, key, names + [name for name in FALLBACK_FONTS if name not in names])
    metadata = spec.get("metadata", {})
    if not isinstance(metadata, dict):
        raise PresentationError("metadata must be an object")
    return DeckManifest(
        schema_version=DECK_SCHEMA_VERSION,
        mode=mode,
        stage=stage,
        template=template,
        assets=assets,
        style=style,
        slides=slides,
        metadata=metadata,
        fonts=fonts,
        asset_cache=asset_cache,
    )


def remove_placeholder_shapes(container: Any, types: Sequence[int]) -> None:
    for shape in list(container.placeholders):
        try:
            placeholder_type = int(shape.placeholder_format.type)
        except Exception:
            continue
        if placeholder_type in types:
            element = shape._element  # noqa: SLF001 - python-pptx has no public removal API
            element.getparent().remove(element)


def strip_footer_placeholders(presentation: Any) -> None:
    """Remove date/footer/slide-number placeholders from the blank layout and its master.

    PowerPoint does not draw them unless a slide instantiates them, but HTML renderers do.
    """
    # MSO_PLACEHOLDER: DATE=16, FOOTER=15, SLIDE_NUMBER=13
    footer_types = (16, 15, 13)
    for layout in presentation.slide_layouts:
        remove_placeholder_shapes(layout, footer_types)
    for master in presentation.slide_masters:
        remove_placeholder_shapes(master, footer_types)


def blank_layout(presentation: Any, preferred_name: Optional[str] = None) -> Any:
    if preferred_name:
        for layout in presentation.slide_layouts:
            if layout.name.strip().lower() == preferred_name.strip().lower():
                return layout
        raise PresentationError("layout '" + preferred_name + "' does not exist in the template")
    for layout in presentation.slide_layouts:
        if layout.name.strip().lower() in ("blank", "空白"):
            return layout
    best = None
    best_count = None
    for layout in presentation.slide_layouts:
        count = len(list(layout.placeholders))
        if best is None or count < best_count:
            best, best_count = layout, count
    if best is None:
        raise PresentationError("the template does not provide any slide layout")
    return best


def remove_all_slides(presentation: Any) -> int:
    """Drop template sample slides while keeping masters, layouts, and theme parts."""
    sldIdLst = presentation.slides._sldIdLst  # noqa: SLF001 - python-pptx has no public API
    removed = 0
    for sldId in list(sldIdLst):
        rId = sldId.rId
        presentation.part.drop_rel(rId)
        sldIdLst.remove(sldId)
        removed += 1
    return removed


def protected_regions_px(presentation: Any, layout: Any) -> List[Dict[str, Any]]:
    """Return non-placeholder master/layout shapes as 1280-canvas rectangles."""
    scale_x = CANVAS_WIDTH / float(presentation.slide_width)
    scale_y = CANVAS_HEIGHT / float(presentation.slide_height)
    regions: List[Dict[str, Any]] = []
    for source, container in (("master", layout.slide_master), ("layout", layout)):
        for shape in container.shapes:
            if getattr(shape, "is_placeholder", False):
                continue
            if shape.width is None or shape.height is None or shape.left is None or shape.top is None:
                continue
            width = float(shape.width) * scale_x
            height = float(shape.height) * scale_y
            if width >= CANVAS_WIDTH * 0.98 and height >= CANVAS_HEIGHT * 0.98:
                continue  # full-bleed background image, not a protected element
            regions.append(
                {
                    "source": source,
                    "name": shape.name,
                    "x": round(float(shape.left) * scale_x, 1),
                    "y": round(float(shape.top) * scale_y, 1),
                    "width": round(width, 1),
                    "height": round(height, 1),
                }
            )
    return regions


def theme_summary(presentation: Any) -> Dict[str, Any]:
    """Extract theme colors and fonts from the first slide master."""
    summary: Dict[str, Any] = {"colors": {}, "fonts": {}}
    try:
        master = presentation.slide_masters[0]
        theme_part = None
        for rel in master.part.rels.values():
            if rel.reltype.endswith("/theme"):
                theme_part = rel.target_part
                break
        if theme_part is None:
            return summary
        root = etree.fromstring(theme_part.blob)
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        scheme = root.find(".//a:clrScheme", ns)
        if scheme is not None:
            for child in scheme:
                name = svg_local_name(child.tag)
                color = child.find("a:srgbClr", ns)
                system = child.find("a:sysClr", ns)
                if color is not None:
                    summary["colors"][name] = color.get("val")
                elif system is not None:
                    summary["colors"][name] = system.get("lastClr") or system.get("val")
        font_scheme = root.find(".//a:fontScheme", ns)
        if font_scheme is not None:
            for role in ("majorFont", "minorFont"):
                node = font_scheme.find("a:" + role, ns)
                if node is None:
                    continue
                latin = node.find("a:latin", ns)
                east_asian = node.find("a:ea", ns)
                summary["fonts"][role] = {
                    "latin": latin.get("typeface") if latin is not None else None,
                    "east_asian": east_asian.get("typeface") if east_asian is not None else None,
                }
    except Exception:
        return summary
    return summary


class SvgSlideConverter:
    def __init__(
        self,
        manifest: DeckManifest,
        measurer: TextMeasurer,
        report: ConversionReport,
        protected: Sequence[Mapping[str, Any]],
        scale: Tuple[float, float],
    ) -> None:
        self.manifest = manifest
        self.measurer = measurer
        self.report = report
        self.protected = list(protected)
        self.scale_x, self.scale_y = scale
        self.gradients: Dict[str, Dict[str, Any]] = {}
        self.vector_assets: Dict[str, Dict[str, Any]] = {}
        self.page = ""
        self.element_count = 0
        self.page_ids: set[str] = set()

    # -- geometry helpers ---------------------------------------------------

    def emu_x(self, value: float) -> Any:
        return Emu(int(round(value * self.scale_x * EMU_PER_PX)))

    def emu_y(self, value: float) -> Any:
        return Emu(int(round(value * self.scale_y * EMU_PER_PX)))

    def record_bounds(self, name: str, x: float, y: float, width: float, height: float) -> None:
        if x < -0.5 or y < -0.5 or x + width > CANVAS_WIDTH + 0.5 or y + height > CANVAS_HEIGHT + 0.5:
            self.report.out_of_bounds.append(
                {"slide": self.page, "element": name, "x": round(x, 1), "y": round(y, 1), "width": round(width, 1), "height": round(height, 1)}
            )
        for region in self.protected:
            if (
                x < region["x"] + region["width"]
                and x + width > region["x"]
                and y < region["y"] + region["height"]
                and y + height > region["y"]
            ):
                self.report.protected_collisions.append(
                    {"slide": self.page, "element": name, "protected": region["name"], "source": region["source"]}
                )

    def count_element(self, tag: str) -> None:
        self.element_count += 1
        if self.element_count > MAX_SVG_ELEMENTS:
            raise PresentationError(self.page + ": SVG contains more than " + str(MAX_SVG_ELEMENTS) + " elements")
        self.report.element_counts[tag] = self.report.element_counts.get(tag, 0) + 1

    # -- fill / line ---------------------------------------------------------

    def apply_fill(self, shape: Any, style: SvgStyle, fill_value: Optional[str]) -> None:
        fill = shape.fill
        if not fill_value:
            fill.background()
            return
        match = re.fullmatch(r"url\(#([A-Za-z0-9_-]+)\)", fill_value)
        if match:
            gradient = self.gradients.get(match.group(1))
            if gradient is None:
                raise PresentationError(self.page + ": unknown gradient " + match.group(1))
            fill.gradient()
            fill.gradient_angle = gradient["angle"]
            self.write_gradient_stops(fill, gradient["stops"])
            return
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(fill_value)
        if style.fill_opacity < 1.0:
            srgb = fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))  # noqa: SLF001
            alpha = etree.SubElement(srgb, qn("a:alpha"))
            alpha.set("val", str(int(round(style.fill_opacity * 100000))))

    def write_gradient_stops(self, fill: Any, stops: Sequence[Tuple[float, str, float]]) -> None:
        gs_lst = fill._fill._gradFill.gsLst  # noqa: SLF001 - python-pptx exposes no stop API
        for gs in list(gs_lst):
            gs_lst.remove(gs)
        for offset, color, opacity in stops:
            gs = etree.SubElement(gs_lst, qn("a:gs"))
            gs.set("pos", str(int(round(offset * 100000))))
            srgb = etree.SubElement(gs, qn("a:srgbClr"))
            srgb.set("val", color)
            if opacity < 1.0:
                alpha = etree.SubElement(srgb, qn("a:alpha"))
                alpha.set("val", str(int(round(opacity * 100000))))

    def apply_line(self, shape: Any, style: SvgStyle) -> None:
        if style.stroke is None:
            shape.line.fill.background()
            return
        shape.line.color.rgb = RGBColor.from_string(style.stroke)
        shape.line.width = px_to_emu(style.stroke_width)

    def add_autoshape(self, slide: Any, kind: Any, x: float, y: float, width: float, height: float, style: SvgStyle, fill_value: Optional[str], name: str) -> Any:
        shape = slide.shapes.add_shape(kind, self.emu_x(x), self.emu_y(y), self.emu_x(width), self.emu_y(height))
        shape.name = name
        shape.shadow.inherit = False
        self.apply_fill(shape, style, fill_value)
        self.apply_line(shape, style)
        self.record_bounds(name, x, y, width, height)
        return shape


    # -- semantic roles --------------------------------------------------------

    SEMANTIC_ROLES = ("title", "subtitle", "text_block", "image", "chart", "table", "footer")

    def tag_role(self, shape: Any, element: Any) -> None:
        """Persist data-role as alt text so the review overlay can classify the shape."""
        role = (element.get("data-role") or "").strip().lower()
        if not role:
            return
        if role not in self.SEMANTIC_ROLES:
            raise PresentationError(
                self.page + ": data-role must be one of " + ", ".join(self.SEMANTIC_ROLES) + ", got " + role
            )
        for nv in ("nvSpPr", "nvPicPr", "nvGraphicFramePr", "nvGrpSpPr", "nvCxnSpPr"):
            container = shape._element.find(qn("p:" + nv))  # noqa: SLF001
            if container is not None:
                container.find(qn("p:cNvPr")).set("descr", "arcforge:role=" + role)
                return
        raise PresentationError(self.page + ": cannot persist data-role on " + shape.name)

    def register_element_id(self, element: Any) -> None:
        """Every explicit id on a page must be unique; the review panel addresses shapes by it."""
        element_id = (element.get("id") or "").strip()
        role = element.get("data-role")
        if role is not None:
            if role.strip().lower() not in self.SEMANTIC_ROLES:
                raise PresentationError(self.page + ": data-role must be one of " + ", ".join(self.SEMANTIC_ROLES))
            if svg_local_name(element.tag) not in ("g", "rect", "circle", "ellipse", "line", "polygon", "path", "text", "image"):
                raise PresentationError(self.page + ": data-role requires a rendered element or group")
            if (element.get("data-arcforge") or "").strip().lower() == "background":
                raise PresentationError(self.page + ": a slide background cannot be a semantic element")
        if not element_id:
            if role is not None:
                raise PresentationError(self.page + ": elements with data-role must also have an id")
            return
        if element.get("id") != element_id or any(character.isspace() for character in element_id):
            raise PresentationError(self.page + ": element ids must not contain whitespace")
        if element_id in self.page_ids:
            raise PresentationError(self.page + ": duplicate element id '" + element_id + "'")
        self.page_ids.add(element_id)

    # -- element handlers ----------------------------------------------------

    def handle_defs(self, defs: Any) -> None:
        for child in defs:
            tag = svg_local_name(child.tag)
            if tag != "linearGradient":
                raise PresentationError(self.page + ": <defs> may only contain <linearGradient>, got <" + tag + ">")
            gradient_id = child.get("id")
            if not gradient_id:
                raise PresentationError(self.page + ": linearGradient requires an id")
            stops: List[Tuple[float, str, float]] = []
            for stop in child:
                if svg_local_name(stop.tag) != "stop":
                    raise PresentationError(self.page + ": linearGradient may only contain <stop>")
                raw_offset = str(stop.get("offset", "0")).strip()
                offset = float(raw_offset.rstrip("%")) / 100.0 if raw_offset.endswith("%") else float(raw_offset)
                color = parse_svg_color(stop.get("stop-color", "#000000"), self.page + " stop-color") or "000000"
                opacity = float(stop.get("stop-opacity", "1"))
                stops.append((max(0.0, min(1.0, offset)), color, max(0.0, min(1.0, opacity))))
            if len(stops) < 2:
                raise PresentationError(self.page + ": linearGradient requires at least two stops")

            def coordinate(name: str, default: str) -> float:
                raw = str(child.get(name, default)).strip()
                return float(raw.rstrip("%")) if raw.endswith("%") else float(raw) * 100.0

            x1, y1 = coordinate("x1", "0%"), coordinate("y1", "0%")
            x2, y2 = coordinate("x2", "100%"), coordinate("y2", "0%")
            angle = math.degrees(math.atan2(y2 - y1, x2 - x1)) % 360.0
            self.gradients[gradient_id] = {"stops": stops, "angle": angle}

    def handle_rect(self, slide: Any, element: Any, style: SvgStyle) -> None:
        name = element.get("id") or "Rect"
        x = parse_svg_length(element.get("x", 0), self.page + " rect.x") + style.dx
        y = parse_svg_length(element.get("y", 0), self.page + " rect.y") + style.dy
        width = parse_svg_length(element.get("width"), self.page + " rect.width")
        height = parse_svg_length(element.get("height"), self.page + " rect.height")
        if width <= 0 or height <= 0:
            raise PresentationError(self.page + ": rect " + name + " must have positive width and height")
        role = (element.get("data-arcforge") or "").strip().lower()
        if role == "background":
            fill_value = style.fill
            if fill_value and fill_value.startswith("url("):
                shape = self.add_autoshape(slide, MSO_SHAPE.RECTANGLE, 0, 0, CANVAS_WIDTH, CANVAS_HEIGHT, style, fill_value, "Background")
                shape.line.fill.background()
            elif fill_value:
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor.from_string(fill_value)
            return
        if role == "chart":
            self.handle_chart(slide, element, name, x, y, width, height)
            return
        if role:
            raise PresentationError(self.page + ": unsupported data-arcforge role '" + role + "'")
        radius = parse_svg_length(element.get("rx", 0), self.page + " rect.rx")
        if radius > 0:
            shape = self.add_autoshape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, width, height, style, style.fill, name)
            shape.adjustments[0] = max(0.0, min(0.5, radius / min(width, height)))
            self.tag_role(shape, element)
        else:
            shape = self.add_autoshape(slide, MSO_SHAPE.RECTANGLE, x, y, width, height, style, style.fill, name)
            self.tag_role(shape, element)

    def handle_ellipse(self, slide: Any, element: Any, style: SvgStyle) -> None:
        name = element.get("id") or "Oval"
        cx = parse_svg_length(element.get("cx", 0), self.page + " cx") + style.dx
        cy = parse_svg_length(element.get("cy", 0), self.page + " cy") + style.dy
        if svg_local_name(element.tag) == "circle":
            rx = ry = parse_svg_length(element.get("r"), self.page + " r")
        else:
            rx = parse_svg_length(element.get("rx"), self.page + " rx")
            ry = parse_svg_length(element.get("ry"), self.page + " ry")
        if rx <= 0 or ry <= 0:
            raise PresentationError(self.page + ": " + name + " must have a positive radius")
        shape = self.add_autoshape(slide, MSO_SHAPE.OVAL, cx - rx, cy - ry, 2 * rx, 2 * ry, style, style.fill, name)
        self.tag_role(shape, element)

    def handle_line(self, slide: Any, element: Any, style: SvgStyle) -> None:
        name = element.get("id") or "Line"
        x1 = parse_svg_length(element.get("x1", 0), self.page + " x1") + style.dx
        y1 = parse_svg_length(element.get("y1", 0), self.page + " y1") + style.dy
        x2 = parse_svg_length(element.get("x2", 0), self.page + " x2") + style.dx
        y2 = parse_svg_length(element.get("y2", 0), self.page + " y2") + style.dy
        connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, self.emu_x(x1), self.emu_y(y1), self.emu_x(x2), self.emu_y(y2))
        connector.name = name
        self.tag_role(connector, element)
        if style.stroke is None:
            connector.line.fill.background()
        else:
            connector.line.color.rgb = RGBColor.from_string(style.stroke)
            connector.line.width = px_to_emu(style.stroke_width)
        self.record_bounds(name, min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1))

    def handle_path(self, slide: Any, element: Any, style: SvgStyle) -> None:
        name = element.get("id") or "Freeform"
        closed = False
        points: List[Tuple[float, float]] = []
        if svg_local_name(element.tag) == "polygon":
            numbers = [float(value) for value in re.findall(r"-?\d+(?:\.\d+)?(?:e-?\d+)?", element.get("points", ""))]
            if len(numbers) < 6 or len(numbers) % 2:
                raise PresentationError(self.page + ": polygon " + name + " needs at least three x,y pairs")
            points = list(zip(numbers[0::2], numbers[1::2]))
            closed = True
        else:
            data = element.get("d", "")
            tokens = re.findall(r"[A-Za-z]|-?\d+(?:\.\d+)?(?:e-?\d+)?", data)
            command = None
            index = 0
            current = (0.0, 0.0)
            while index < len(tokens):
                token = tokens[index]
                if token.isalpha():
                    if token in ("Z", "z"):
                        closed = True
                        index += 1
                        continue
                    if token not in ("M", "L", "m", "l", "H", "V", "h", "v"):
                        raise PresentationError(
                            self.page + ": path " + name + " uses unsupported command '" + token + "'; only M, L, H, V and Z are allowed"
                        )
                    command = token
                    index += 1
                    continue
                if command is None:
                    raise PresentationError(self.page + ": path " + name + " must start with M")
                if command in ("H", "h", "V", "v"):
                    value = float(token)
                    index += 1
                    if command == "H":
                        current = (value, current[1])
                    elif command == "h":
                        current = (current[0] + value, current[1])
                    elif command == "V":
                        current = (current[0], value)
                    else:
                        current = (current[0], current[1] + value)
                    points.append(current)
                    continue
                if index + 1 >= len(tokens):
                    raise PresentationError(self.page + ": path " + name + " has a dangling coordinate")
                px, py = float(token), float(tokens[index + 1])
                index += 2
                if command in ("m", "l"):
                    current = (current[0] + px, current[1] + py)
                else:
                    current = (px, py)
                points.append(current)
                if command == "M":
                    command = "L"
                elif command == "m":
                    command = "l"
            if len(points) < 2:
                raise PresentationError(self.page + ": path " + name + " needs at least two points")
        points = [(px + style.dx, py + style.dy) for px, py in points]
        xs = [point[0] for point in points]
        ys = [point[1] for point in points]
        builder = slide.shapes.build_freeform(self.emu_x(points[0][0]), self.emu_y(points[0][1]), scale=1.0)
        builder.add_line_segments([(self.emu_x(px), self.emu_y(py)) for px, py in points[1:]], close=closed)
        shape = builder.convert_to_shape()
        shape.name = name
        shape.shadow.inherit = False
        self.tag_role(shape, element)
        self.apply_fill(shape, style, style.fill if closed else None)
        self.apply_line(shape, style)
        self.record_bounds(name, min(xs), min(ys), max(xs) - min(xs), max(ys) - min(ys))

    def apply_font(self, run: Any, family: str, size_px: float, bold: bool, color: str) -> None:
        font = run.font
        font.size = Pt(size_px * 0.75)
        font.bold = bold
        font.color.rgb = RGBColor.from_string(color)
        font.name = family
        rPr = run._r.get_or_add_rPr()  # noqa: SLF001 - python-pptx only writes a:latin
        for tag in ("a:latin", "a:ea", "a:cs"):
            node = rPr.find(qn(tag))
            if node is None:
                node = etree.SubElement(rPr, qn(tag))
            node.set("typeface", family)

    def handle_text(self, slide: Any, element: Any, style: SvgStyle) -> None:
        name = element.get("id") or "Text"
        x = parse_svg_length(element.get("x", 0), self.page + " text.x") + style.dx
        baseline = parse_svg_length(element.get("y", 0), self.page + " text.y") + style.dy
        size = style.font_size
        bold = style.font_weight in ("bold", "bolder", "600", "700", "800", "900")
        family = style.font_family
        lines: List[Tuple[str, float]] = []
        spans = [child for child in element if svg_local_name(child.tag) == "tspan"]
        for child in element:
            if svg_local_name(child.tag) != "tspan":
                raise PresentationError(self.page + ": <text> " + name + " may only contain <tspan>")
        if spans:
            leading = (element.text or "").strip()
            if leading:
                lines.append((leading, baseline))
            current = baseline
            for span in spans:
                if len(span):
                    raise PresentationError(self.page + ": nested <tspan> is not supported in " + name)
                if span.get("y") is not None:
                    current = parse_svg_length(span.get("y"), self.page + " tspan.y") + style.dy
                elif span.get("dy") is not None:
                    current += parse_svg_length(span.get("dy"), self.page + " tspan.dy")
                elif lines:
                    current += size * 1.2
                lines.append((" ".join((span.text or "").split()), current))
        else:
            lines.append((" ".join((element.text or "").split()), baseline))
        lines = [(text, y) for text, y in lines if text]
        if not lines:
            return
        declared_width = element.get("data-width")
        width_limit = parse_svg_length(declared_width, self.page + " data-width") if declared_width else None
        wrapped: List[str] = []
        for text, _ in lines:
            if width_limit:
                wrapped.extend(self.measurer.wrap(text, family, bold, size, width_limit))
            else:
                wrapped.append(text)
        widest = max(self.measurer.width(text, family, bold, size) for text in wrapped)
        if width_limit and len(wrapped) > len(lines):
            self.report.text_overflows.append(
                {"slide": self.page, "element": name, "declared_width": width_limit, "needed_width": round(widest, 1), "wrapped_lines": len(wrapped), "declared_lines": len(lines)}
            )
        line_gap = (lines[1][1] - lines[0][1]) if len(lines) > 1 else size * 1.2
        if line_gap <= 0:
            line_gap = size * 1.2
        box_width = width_limit if width_limit else widest + size * 0.4
        ascent = size * 0.88
        top = lines[0][1] - ascent
        height = line_gap * (len(wrapped) - 1) + size * 1.25
        if style.text_anchor == "middle":
            left = x - box_width / 2
            alignment = PP_ALIGN.CENTER
        elif style.text_anchor == "end":
            left = x - box_width
            alignment = PP_ALIGN.RIGHT
        else:
            left = x
            alignment = PP_ALIGN.LEFT
        color = style.fill if style.fill and not style.fill.startswith("url(") else "000000"
        box = slide.shapes.add_textbox(self.emu_x(left), self.emu_y(top), self.emu_x(box_width), self.emu_y(height))
        box.name = name
        self.tag_role(box, element)
        frame = box.text_frame
        frame.word_wrap = bool(width_limit)
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for index, text in enumerate(wrapped):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.alignment = alignment
            paragraph.line_spacing = Pt(line_gap * 0.75)
            run = paragraph.add_run()
            run.text = text
            self.apply_font(run, family, size, bold, color)
        self.record_bounds(name, left, top, box_width, height)

    def handle_image(self, slide: Any, element: Any, style: SvgStyle) -> None:
        name = element.get("id") or "Picture"
        for forbidden in ("href", XLINK_NS + "href"):
            if element.get(forbidden) is not None:
                raise PresentationError(self.page + ": <image> must reference an asset with data-asset, not href")
        asset_id = (element.get("data-asset") or "").strip()
        if not asset_id:
            raise PresentationError(self.page + ": <image> " + name + " requires data-asset")
        path = self.manifest.assets.get(asset_id)
        if path is None:
            raise PresentationError(self.page + ": data-asset '" + asset_id + "' is not declared in the manifest assets")
        x = parse_svg_length(element.get("x", 0), self.page + " image.x") + style.dx
        y = parse_svg_length(element.get("y", 0), self.page + " image.y") + style.dy
        width = parse_svg_length(element.get("width"), self.page + " image.width")
        height = parse_svg_length(element.get("height"), self.page + " image.height")
        if width <= 0 or height <= 0:
            raise PresentationError(self.page + ": image " + name + " must have positive width and height")
        mode = (element.get("preserveAspectRatio") or "xMidYMid meet").strip().lower()
        if mode == "none":
            raise PresentationError(self.page + ": image " + name + " must keep its aspect ratio (meet or slice)")
        if path.suffix.lower() == ".svg":
            self.handle_vector_asset(slide, element, name, asset_id, path, x, y, width, height, mode)
            return
        for attribute in ("data-fill", "data-stroke", "data-render"):
            if element.get(attribute) is not None:
                raise PresentationError(
                    self.page + ": " + attribute + " on image " + name + " is only valid for SVG assets"
                )
        self.place_picture(slide, name, path, x, y, width, height, mode, element)

    def place_picture(
        self,
        slide: Any,
        name: str,
        path: Path,
        x: float,
        y: float,
        width: float,
        height: float,
        mode: str,
        element: Any = None,
    ) -> None:
        with Image.open(path) as image:
            image_width, image_height = image.size
        if image_width <= 0 or image_height <= 0:
            raise PresentationError("Image has invalid dimensions: " + str(path))
        image_ratio = image_width / image_height
        frame_ratio = width / height
        if "slice" in mode:
            picture = slide.shapes.add_picture(str(path), self.emu_x(x), self.emu_y(y), self.emu_x(width), self.emu_y(height))
            if image_ratio > frame_ratio:
                crop = (1.0 - frame_ratio / image_ratio) / 2.0
                picture.crop_left = picture.crop_right = crop
            elif image_ratio < frame_ratio:
                crop = (1.0 - image_ratio / frame_ratio) / 2.0
                picture.crop_top = picture.crop_bottom = crop
            bounds = (x, y, width, height)
        else:
            if image_ratio >= frame_ratio:
                fitted_width, fitted_height = width, width / image_ratio
            else:
                fitted_width, fitted_height = height * image_ratio, height
            left = x + (width - fitted_width) / 2
            top = y + (height - fitted_height) / 2
            picture = slide.shapes.add_picture(str(path), self.emu_x(left), self.emu_y(top), self.emu_x(fitted_width), self.emu_y(fitted_height))
            bounds = (left, top, fitted_width, fitted_height)
        picture.name = name
        self.record_bounds(name, *bounds)
        if element is not None:
            self.tag_role(picture, element)

    # -- vector assets (SVG normalized by ArcForge) ----------------------------

    def load_vector_asset(self, asset_id: str, source: Path) -> Dict[str, Any]:
        cached = self.vector_assets.get(asset_id)
        if cached is not None:
            return cached
        cache_root = self.manifest.asset_cache
        if cache_root is None:
            raise PresentationError(
                self.page + ": asset '" + asset_id + "' is an SVG; run create or validate through the OfficeRuntime tool so ArcForge can normalize it"
            )
        entry = cache_root / asset_id / "shapes.json"
        if not entry.is_file():
            raise PresentationError(
                self.page + ": asset '" + asset_id + "' has no prepared entry in " + str(cache_root) + "; run validate again"
            )
        try:
            data = json.loads(entry.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as error:
            raise PresentationError(self.page + ": asset '" + asset_id + "' cache is unreadable: " + str(error)) from error
        if not isinstance(data, dict) or data.get("schema") != VECTOR_ASSET_SCHEMA:
            raise PresentationError(self.page + ": asset '" + asset_id + "' cache uses an unsupported schema")
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        if data.get("source_sha256") != digest:
            raise PresentationError(
                self.page + ": asset '" + asset_id + "' changed after it was prepared; run validate again"
            )
        data["_directory"] = entry.parent
        self.vector_assets[asset_id] = data
        return data

    def handle_vector_asset(
        self,
        slide: Any,
        element: Any,
        name: str,
        asset_id: str,
        source: Path,
        x: float,
        y: float,
        width: float,
        height: float,
        mode: str,
    ) -> None:
        asset = self.load_vector_asset(asset_id, source)
        render = (element.get("data-render") or "auto").strip().lower()
        if render not in ("auto", "shapes", "raster"):
            raise PresentationError(self.page + ": data-render on " + name + " must be auto, shapes, or raster")
        fill_override = parse_svg_color(element.get("data-fill"), self.page + " data-fill")
        stroke_override = parse_svg_color(element.get("data-stroke"), self.page + " data-stroke")
        if stroke_override is None:
            # data-fill is "the icon color": monochrome stroke icons pick it up too unless
            # data-stroke says otherwise.
            stroke_override = fill_override
        if self.manifest.stage == "plan":
            fill_override = PLAN_STAGE_GRAY
            stroke_override = PLAN_STAGE_GRAY
        reasons = [str(reason) for reason in asset.get("reasons", [])]
        if "slice" in mode:
            reasons = reasons + ["preserveAspectRatio slice needs a picture"]
        native_possible = asset.get("mode") == "shapes" and "slice" not in mode
        if render == "shapes" and not native_possible:
            raise PresentationError(
                self.page + ": asset '" + asset_id + "' cannot become native shapes (" + ", ".join(reasons) + "); remove data-render=\"shapes\" or simplify the SVG"
            )
        use_raster = render == "raster" or not native_possible
        directory: Path = asset["_directory"]
        record: Dict[str, Any] = {"slide": self.page, "element": name, "asset": asset_id}
        if use_raster:
            raster = directory / str(asset.get("raster") or "raster.png")
            if not raster.is_file():
                raise PresentationError(self.page + ": asset '" + asset_id + "' has no raster fallback; run validate again")
            if self.manifest.stage == "plan":
                raster = self.grayscale_copy(raster)
            self.place_picture(slide, name, raster, x, y, width, height, mode, element)
            record.update({"mode": "raster", "reasons": reasons if render != "raster" else ["data-render=raster"]})
            if render != "raster" and reasons:
                self.report.warnings.append(
                    self.page + ": asset '" + asset_id + "' was placed as a picture because of " + ", ".join(reasons)
                )
        else:
            count = self.place_vector_shapes(slide, name, asset, x, y, width, height, fill_override, stroke_override, element)
            record.update({"mode": "shapes", "shapes": count})
        self.report.asset_renders.append(record)
        self.record_bounds(name, x, y, width, height)

    def grayscale_copy(self, raster: Path) -> Path:
        target = raster.with_name(raster.stem + "-plan.png")
        if not target.is_file() or target.stat().st_mtime < raster.stat().st_mtime:
            with Image.open(raster) as image:
                image.convert("LA").save(target, format="PNG")
        return target

    def place_vector_shapes(
        self,
        slide: Any,
        name: str,
        asset: Mapping[str, Any],
        x: float,
        y: float,
        width: float,
        height: float,
        fill_override: Optional[str],
        stroke_override: Optional[str],
        element: Any = None,
    ) -> int:
        view_width = float(asset.get("width") or 0)
        view_height = float(asset.get("height") or 0)
        if view_width <= 0 or view_height <= 0:
            raise PresentationError(self.page + ": asset '" + str(asset.get("asset_id")) + "' has an empty view box")
        scale = min(width / view_width, height / view_height)
        origin_x = x + (width - view_width * scale) / 2.0
        origin_y = y + (height - view_height * scale) / 2.0
        group = slide.shapes.add_group_shape()
        group.name = name
        if element is not None:
            self.tag_role(group, element)
        group_tree = group.shapes._spTree  # noqa: SLF001 - python-pptx has no public curve API
        count = 0
        for index, spec in enumerate(asset.get("paths", []), start=1):
            fill_spec = spec.get("fill")
            stroke_spec = spec.get("stroke")
            frame = spec.get("stroke_bbox") if stroke_spec else spec.get("bbox")
            if not isinstance(frame, list) or len(frame) != 4:
                raise PresentationError(self.page + ": vector path " + str(index) + " has no bounding box")
            frame_x, frame_y, frame_w, frame_h = (float(value) for value in frame)
            left = origin_x + frame_x * scale
            top = origin_y + frame_y * scale
            extent_x = max(1, int(round(frame_w * scale * self.scale_x * EMU_PER_PX)))
            extent_y = max(1, int(round(frame_h * scale * self.scale_y * EMU_PER_PX)))
            sp = group_tree.add_freeform_sp(self.emu_x(left), self.emu_y(top), Emu(extent_x), Emu(extent_y))
            path = sp.add_path(w=Emu(extent_x), h=Emu(extent_y))

            def local_x(value: float) -> int:
                return int(round((origin_x + value * scale - left) * self.scale_x * EMU_PER_PX))

            def local_y(value: float) -> int:
                return int(round((origin_y + value * scale - top) * self.scale_y * EMU_PER_PX))

            for segment in spec.get("segments", []):
                command = segment[0]
                if command == "M":
                    path.add_moveTo(Emu(local_x(segment[1])), Emu(local_y(segment[2])))
                elif command == "L":
                    path.add_lnTo(Emu(local_x(segment[1])), Emu(local_y(segment[2])))
                elif command == "Q":
                    node = etree.SubElement(path, qn("a:quadBezTo"))
                    for px, py in ((segment[1], segment[2]), (segment[3], segment[4])):
                        point = etree.SubElement(node, qn("a:pt"))
                        point.set("x", str(local_x(px)))
                        point.set("y", str(local_y(py)))
                elif command == "C":
                    node = etree.SubElement(path, qn("a:cubicBezTo"))
                    for px, py in ((segment[1], segment[2]), (segment[3], segment[4]), (segment[5], segment[6])):
                        point = etree.SubElement(node, qn("a:pt"))
                        point.set("x", str(local_x(px)))
                        point.set("y", str(local_y(py)))
                elif command == "Z":
                    path.add_close()
                else:
                    raise PresentationError(self.page + ": vector path uses unknown command " + str(command))
            shape = group.shapes._shape_factory(sp)  # noqa: SLF001
            shape.name = name + "/" + str(index)
            shape.shadow.inherit = False
            self.apply_vector_fill(shape, fill_spec, fill_override)
            self.apply_vector_stroke(shape, stroke_spec, scale, stroke_override)
            count += 1
        if count == 0:
            group._element.getparent().remove(group._element)  # noqa: SLF001
            self.report.warnings.append(self.page + ": vector asset " + name + " produced no shapes")
            return 0
        group._element.recalculate_extents()  # noqa: SLF001
        self.report.element_counts["vector_shapes"] = self.report.element_counts.get("vector_shapes", 0) + count
        return count

    def apply_vector_fill(self, shape: Any, spec: Optional[Mapping[str, Any]], override: Optional[str]) -> None:
        fill = shape.fill
        if not spec:
            fill.background()
            return
        opacity = max(0.0, min(1.0, float(spec.get("opacity", 1.0))))
        if spec.get("type") == "linear" and override is None:
            stops = [
                (
                    max(0.0, min(1.0, float(stop.get("offset", 0.0)))),
                    normalize_color(str(stop.get("color", "000000")), "gradient stop"),
                    max(0.0, min(1.0, float(stop.get("opacity", 1.0)))),
                )
                for stop in spec.get("stops", [])
            ]
            if len(stops) >= 2:
                fill.gradient()
                fill.gradient_angle = float(spec.get("angle", 0.0))
                self.write_gradient_stops(fill, stops)
                return
            color = stops[0][1] if stops else "000000"
        else:
            color = override or normalize_color(str(spec.get("color", "000000")), "vector fill")
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(color)
        if opacity < 1.0:
            srgb = fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))  # noqa: SLF001
            alpha = etree.SubElement(srgb, qn("a:alpha"))
            alpha.set("val", str(int(round(opacity * 100000))))

    def apply_vector_stroke(
        self, shape: Any, spec: Optional[Mapping[str, Any]], scale: float, override: Optional[str]
    ) -> None:
        line = shape.line
        if not spec:
            line.fill.background()
            return
        color = override or normalize_color(str(spec.get("color", "000000")), "vector stroke")
        line.color.rgb = RGBColor.from_string(color)
        width_px = float(spec.get("width", 1.0)) * scale
        line.width = Emu(max(1, int(round(width_px * self.scale_x * EMU_PER_PX))))
        ln = line._get_or_add_ln()  # noqa: SLF001 - caps and joins have no public API
        ln.set("cap", {"round": "rnd", "square": "sq"}.get(str(spec.get("linecap", "butt")), "flat"))
        opacity = max(0.0, min(1.0, float(spec.get("opacity", 1.0))))
        if opacity < 1.0:
            srgb = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            alpha = etree.SubElement(srgb, qn("a:alpha"))
            alpha.set("val", str(int(round(opacity * 100000))))
        if spec.get("dash"):
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
        join = str(spec.get("linejoin", "miter"))
        if join == "round":
            etree.SubElement(ln, qn("a:round"))
        elif join == "bevel":
            etree.SubElement(ln, qn("a:bevel"))
        else:
            etree.SubElement(ln, qn("a:miter")).set("lim", "800000")

    def handle_chart(self, slide: Any, element: Any, name: str, x: float, y: float, width: float, height: float) -> None:
        raw = element.get("data-chart")
        if not raw:
            raise PresentationError(self.page + ": chart " + name + " requires data-chart JSON")
        try:
            spec = json.loads(raw)
        except json.JSONDecodeError as error:
            raise PresentationError(self.page + ": chart " + name + " data-chart is not valid JSON: " + str(error)) from error
        if not isinstance(spec, dict):
            raise PresentationError(self.page + ": chart " + name + " data-chart must be an object")
        chart_kind = str(spec.get("type", "column")).strip().lower()
        if chart_kind not in CHART_TYPES_V3:
            raise PresentationError(self.page + ": chart " + name + " type must be one of " + ", ".join(sorted(CHART_TYPES_V3)))
        categories = spec.get("categories")
        series_list = spec.get("series")
        if not isinstance(categories, list) or not categories:
            raise PresentationError(self.page + ": chart " + name + " requires categories")
        if not isinstance(series_list, list) or not series_list:
            raise PresentationError(self.page + ": chart " + name + " requires series")
        data = ChartData()
        data.categories = [str(item) for item in categories]
        for series in series_list:
            if not isinstance(series, dict) or "values" not in series:
                raise PresentationError(self.page + ": chart " + name + " series entries need name and values")
            values = series["values"]
            if not isinstance(values, list) or len(values) != len(categories):
                raise PresentationError(self.page + ": chart " + name + " series values must match the category count")
            try:
                numbers = [float(value) for value in values]
            except (TypeError, ValueError) as error:
                raise PresentationError(self.page + ": chart " + name + " series values must be numbers") from error
            data.add_series(str(series.get("name", "Series")), numbers)
        graphic = slide.shapes.add_chart(chart_type_value(chart_kind), self.emu_x(x), self.emu_y(y), self.emu_x(width), self.emu_y(height), data)
        graphic.name = name
        self.tag_role(graphic, element)
        chart = graphic.chart
        chart.has_legend = len(series_list) > 1 or chart_kind in ("pie", "doughnut")
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP if chart_kind not in ("pie", "doughnut") else XL_LEGEND_POSITION.RIGHT
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(10)
        chart.font.name = self.manifest.fonts.body[0]
        chart.font.size = Pt(10)
        colors = spec.get("colors") or []
        if chart_kind in ("pie", "doughnut"):
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.show_percentage = True
            plot.data_labels.show_legend_key = False
        else:
            for index, series in enumerate(chart.series):
                color = colors[index] if index < len(colors) else None
                if color:
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = RGBColor.from_string(parse_svg_color(str(color), self.page + " chart color") or "000000")
            chart.value_axis.has_major_gridlines = True
            chart.value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string("E5E9F0")
            chart.value_axis.format.line.fill.background()
        self.record_bounds(name, x, y, width, height)

    # -- traversal -----------------------------------------------------------

    def walk(self, slide: Any, element: Any, inherited: SvgStyle) -> None:
        tag = svg_local_name(element.tag)
        if not element.tag.startswith(SVG_NS) and element.tag != tag:
            raise PresentationError(self.page + ": foreign namespace element <" + element.tag + "> is not supported")
        if tag not in SUPPORTED_SVG_ELEMENTS:
            raise PresentationError(self.page + ": unsupported SVG element <" + tag + ">")
        self.count_element(tag)
        style = inherited.inherit(element, self.page)
        if tag == "defs":
            self.handle_defs(element)
        elif tag in ("svg", "g"):
            first_shape = len(slide.shapes)
            for child in element:
                self.walk(slide, child, style)
            if tag == "g" and element.get("data-role") is not None:
                children = list(slide.shapes)[first_shape:]
                if not children:
                    raise PresentationError(self.page + ": semantic group '" + element.get("id") + "' has no rendered elements")
                group = slide.shapes.add_group_shape(children)
                group.name = element.get("id")
                self.tag_role(group, element)
        elif tag == "rect":
            self.handle_rect(slide, element, style)
        elif tag in ("circle", "ellipse"):
            self.handle_ellipse(slide, element, style)
        elif tag == "line":
            self.handle_line(slide, element, style)
        elif tag in ("path", "polygon"):
            self.handle_path(slide, element, style)
        elif tag == "text":
            self.handle_text(slide, element, style)
        elif tag == "image":
            self.handle_image(slide, element, style)
        elif tag in ("title", "desc", "linearGradient", "stop"):
            return

    def convert(self, slide: Any, svg_path: Path, slide_id: str) -> None:
        self.page = slide_id
        self.gradients = {}
        self.element_count = 0
        self.page_ids = set()
        if svg_path.stat().st_size > MAX_SVG_BYTES:
            raise PresentationError(slide_id + ": SVG exceeds " + str(MAX_SVG_BYTES // 1024) + " KiB")
        raw = svg_path.read_bytes()
        if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
            raise PresentationError(slide_id + ": DOCTYPE and entity declarations are not allowed")
        try:
            root = ET.fromstring(raw)
        except ET.ParseError as error:
            raise PresentationError(slide_id + ": SVG is not well-formed XML: " + str(error)) from error
        if svg_local_name(root.tag) != "svg":
            raise PresentationError(slide_id + ": root element must be <svg>")
        view_box = [float(value) for value in re.split(r"[\s,]+", (root.get("viewBox") or "").strip()) if value]
        if view_box != [0.0, 0.0, CANVAS_WIDTH, CANVAS_HEIGHT]:
            raise PresentationError(slide_id + ": viewBox must be exactly '0 0 1280 720'")
        # Validate the entire source tree, including groups, gradients, and tspan nodes
        # that the renderer handles without going through walk().
        for element in root.iter():
            self.register_element_id(element)
        self.walk(slide, root, SvgStyle(fill="000000"))
        if self.manifest.stage == "plan":
            self.check_plan_stage(root, slide_id)

    def check_plan_stage(self, root: Any, slide_id: str) -> None:
        for element in root.iter():
            tag = svg_local_name(element.tag)
            if tag == "linearGradient":
                raise PresentationError(slide_id + ": stage=plan pages must not use gradients")
            for attribute in ("fill", "stroke", "stop-color"):
                value = element.get(attribute)
                if not value or value.strip() in ("none", "transparent") or value.startswith("url("):
                    continue
                color = parse_svg_color(value, slide_id + " " + attribute)
                if color is None:
                    continue
                r, g, b = int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16)
                if max(r, g, b) - min(r, g, b) > 24:
                    raise PresentationError(
                        slide_id + ": stage=plan pages must use grayscale colors only, found #" + color + " on <" + tag + ">"
                    )


def open_deck_base(manifest: DeckManifest) -> Tuple[Any, Any, List[Dict[str, Any]], int]:
    if manifest.template is not None:
        try:
            presentation = Presentation(str(manifest.template))
        except Exception as error:
            raise PresentationError("Failed to open template presentation: " + str(error)) from error
        removed = remove_all_slides(presentation)
        layout = blank_layout(presentation, str(manifest.style.get("layout") or "") or None)
        protected = protected_regions_px(presentation, layout)
        return presentation, layout, protected, removed
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH)
    presentation.slide_height = Inches(SLIDE_HEIGHT)
    strip_footer_placeholders(presentation)
    layout = presentation.slide_layouts[6]
    return presentation, layout, [], 0


def fix_chart_axis_ids(presentation: Any) -> None:
    """python-pptx emits random axId values that may be negative; OOXML requires UInt32."""
    counter = 100_000_000
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False):
                continue
            chart_space = shape.chart._chartSpace  # noqa: SLF001
            mapping: Dict[str, str] = {}
            for node in chart_space.iter(qn("c:axId"), qn("c:crossAx")):
                old = node.get("val")
                if old not in mapping:
                    counter += 1
                    mapping[old] = str(counter)
                node.set("val", mapping[old])


def create_svg_deck(
    spec: Mapping[str, Any], spec_dir: Path, asset_cache: Optional[Path] = None
) -> Tuple[Any, Dict[str, Any]]:
    manifest = load_deck_manifest(spec, spec_dir, asset_cache)
    presentation, layout, protected, removed = open_deck_base(manifest)
    apply_metadata(presentation, {"metadata": manifest.metadata})
    scale = (
        float(presentation.slide_width) / (CANVAS_WIDTH * EMU_PER_PX),
        float(presentation.slide_height) / (CANVAS_HEIGHT * EMU_PER_PX),
    )
    measurer = TextMeasurer()
    report = ConversionReport()
    converter = SvgSlideConverter(manifest, measurer, report, protected, scale)
    slide_ids: List[str] = []
    for entry in manifest.slides:
        slide = presentation.slides.add_slide(layout)
        converter.convert(slide, entry["svg"], entry["slide_id"])
        slide.name = entry["slide_id"]
        if entry.get("notes"):
            apply_notes(slide, entry["notes"])
        slide_ids.append(entry["slide_id"])
    fix_chart_axis_ids(presentation)
    report.missing_fonts = list(measurer.missing_fonts)
    report.measured_with_fonts = measurer.measured_with_fonts
    fingerprints = deck_slide_fingerprints(manifest)
    previous = read_build_stamp(spec_dir)
    changed_slide_ids = [
        slide_id for slide_id in slide_ids if previous.get(slide_id) != fingerprints.get(slide_id)
    ]
    removed_slide_ids = [slide_id for slide_id in previous if slide_id not in fingerprints]
    summary = {
        "schema_version": DECK_SCHEMA_VERSION,
        "mode": manifest.mode,
        "stage": manifest.stage,
        "template": str(manifest.template) if manifest.template else None,
        "template_slides_removed": removed,
        "layout": layout.name,
        "canvas": {"width": CANVAS_WIDTH, "height": CANVAS_HEIGHT},
        "slide_ids": slide_ids,
        "protected_regions": protected,
        "element_counts": report.element_counts,
        "text_overflows": report.text_overflows,
        "out_of_bounds": report.out_of_bounds,
        "protected_collisions": report.protected_collisions,
        "missing_fonts": report.missing_fonts,
        "measured_with_fonts": report.measured_with_fonts,
        "asset_renders": report.asset_renders,
        "changed_slide_ids": changed_slide_ids,
        "removed_slide_ids": removed_slide_ids,
        "slide_fingerprints": fingerprints,
        "warnings": report.warnings,
    }
    return presentation, summary


BUILD_STAMP_NAME = ".arcforge-build.json"


def deck_slide_fingerprints(manifest: DeckManifest) -> Dict[str, str]:
    """One hash per page covering its SVG, notes, and the deck-level inputs that change rendering."""
    shared = hashlib.sha256()
    shared.update(manifest.mode.encode("utf-8"))
    shared.update(manifest.stage.encode("utf-8"))
    if manifest.template is not None:
        try:
            shared.update(manifest.template.read_bytes())
        except OSError:
            shared.update(str(manifest.template).encode("utf-8"))
    for asset_id in sorted(manifest.assets):
        try:
            shared.update(asset_id.encode("utf-8"))
            shared.update(manifest.assets[asset_id].read_bytes())
        except OSError:
            continue
    shared.update(json.dumps(manifest.style, sort_keys=True, ensure_ascii=False).encode("utf-8"))
    shared_digest = shared.digest()
    fingerprints: Dict[str, str] = {}
    for entry in manifest.slides:
        digest = hashlib.sha256()
        digest.update(shared_digest)
        try:
            digest.update(Path(entry["svg"]).read_bytes())
        except OSError:
            digest.update(str(entry["svg"]).encode("utf-8"))
        digest.update(json.dumps(entry.get("notes") or "", ensure_ascii=False).encode("utf-8"))
        digest.update(json.dumps(entry.get("layout") or "", ensure_ascii=False).encode("utf-8"))
        fingerprints[entry["slide_id"]] = digest.hexdigest()
    return fingerprints


def read_build_stamp(spec_dir: Path) -> Dict[str, str]:
    stamp = spec_dir / BUILD_STAMP_NAME
    try:
        data = json.loads(stamp.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    slides = data.get("slides") if isinstance(data, dict) else None
    if not isinstance(slides, dict):
        return {}
    return {str(key): str(value) for key, value in slides.items()}


def write_build_stamp(spec_dir: Path, fingerprints: Mapping[str, str], output_path: Path) -> None:
    stamp = spec_dir / BUILD_STAMP_NAME
    payload = {"schema": 1, "output": str(output_path), "slides": dict(fingerprints)}
    try:
        stamp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass


def is_svg_deck_spec(spec: Mapping[str, Any]) -> bool:
    try:
        return int(spec.get("schema_version", 1)) >= DECK_SCHEMA_VERSION
    except (TypeError, ValueError):
        return False


def normalized_output_path(path_value: str, suffix: str) -> Path:
    path = Path(path_value).expanduser().resolve()
    if path.suffix.lower() != suffix:
        raise PresentationError("Output path must end with " + suffix)
    return path


def atomic_save(presentation: Any, output_path: Path, force: bool) -> None:
    if output_path.exists() and not force:
        raise PresentationError(
            "Output already exists. Use a new path or pass --force only after explicit approval: "
            + str(output_path)
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary_name = tempfile.mkstemp(
        prefix=".arcforge-deck-",
        suffix=".pptx",
        dir=str(output_path.parent),
    )
    os.close(handle)
    temporary_path = Path(temporary_name)
    try:
        presentation.save(temporary_path)
        Presentation(temporary_path)
        os.replace(temporary_path, output_path)
    except Exception:
        try:
            temporary_path.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def shape_is_out_of_bounds(shape: Any, slide_width: int, slide_height: int) -> bool:
    return (
        shape.left < 0
        or shape.top < 0
        or shape.left + shape.width > slide_width
        or shape.top + shape.height > slide_height
    )


def slide_title(slide: Any) -> str:
    for shape in slide.shapes:
        if shape.name == "ArcForge Title" and getattr(shape, "has_text_frame", False):
            return shape.text.strip()
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip():
            return shape.text.strip().splitlines()[0]
    return ""


def inspect_presentation(path: Path) -> Dict[str, Any]:
    if not path.is_file():
        raise PresentationError("Presentation does not exist: " + str(path))
    try:
        presentation = Presentation(path)
    except Exception as error:
        raise PresentationError(
            "Failed to inspect presentation: " + str(error)
        ) from error

    slides: List[Dict[str, Any]] = []
    total_images = 0
    total_tables = 0
    total_charts = 0
    total_out_of_bounds = 0
    for index, slide in enumerate(presentation.slides, start=1):
        images = 0
        tables = 0
        charts = 0
        out_of_bounds = 0
        text_characters = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images += 1
            if getattr(shape, "has_table", False):
                tables += 1
            if getattr(shape, "has_chart", False):
                charts += 1
            if getattr(shape, "has_text_frame", False):
                text_characters += len(shape.text)
            if shape_is_out_of_bounds(
                shape, presentation.slide_width, presentation.slide_height
            ):
                out_of_bounds += 1
        title = slide_title(slide)
        total_images += images
        total_tables += tables
        total_charts += charts
        total_out_of_bounds += out_of_bounds
        notes = ""
        if slide.has_notes_slide:
            try:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            except Exception:
                notes = ""
        slides.append(
            {
                "number": index,
                "title": title,
                "shape_count": len(slide.shapes),
                "text_characters": text_characters,
                "image_count": images,
                "table_count": tables,
                "chart_count": charts,
                "out_of_bounds_shapes": out_of_bounds,
                "has_notes": bool(notes),
            }
        )

    layouts: List[Dict[str, Any]] = []
    for layout_index, layout in enumerate(presentation.slide_layouts):
        placeholders = []
        for shape in layout.placeholders:
            try:
                placeholders.append(str(shape.placeholder_format.type).split(".")[-1].split(" ")[0])
            except Exception:
                placeholders.append("unknown")
        layouts.append({"index": layout_index, "name": layout.name, "placeholders": placeholders})
    try:
        base_layout = blank_layout(presentation)
        protected = protected_regions_px(presentation, base_layout)
        base_layout_name = base_layout.name
    except PresentationError:
        protected = []
        base_layout_name = None

    return {
        "path": str(path.resolve()),
        "size_bytes": path.stat().st_size,
        "slide_count": len(slides),
        "slide_size_inches": {
            "width": round(presentation.slide_width / Inches(1), 3),
            "height": round(presentation.slide_height / Inches(1), 3),
        },
        "slides": slides,
        "total_images": total_images,
        "total_tables": total_tables,
        "total_charts": total_charts,
        "total_out_of_bounds_shapes": total_out_of_bounds,
        "missing_title_slides": [
            item["number"] for item in slides if not item["title"]
        ],
        "template": {
            "canvas": {"width": CANVAS_WIDTH, "height": CANVAS_HEIGHT},
            "aspect_ratio": round(float(presentation.slide_width) / float(presentation.slide_height), 4),
            "layouts": layouts,
            "base_layout": base_layout_name,
            "protected_regions": protected,
            "theme": theme_summary(presentation),
        },
        "visually_rendered": False,
    }


def soffice_candidates() -> List[Path]:
    candidates: List[Path] = []
    for key in ("ARCFORGE_SOFFICE_PATH", "LIVEAGENT_SOFFICE_PATH"):
        raw = os.environ.get(key)
        if raw:
            candidates.append(Path(raw).expanduser())
    for command in ("soffice", "libreoffice"):
        resolved = shutil.which(command)
        if resolved:
            candidates.append(Path(resolved))
    for key in ("ProgramFiles", "ProgramFiles(x86)", "LOCALAPPDATA"):
        root = os.environ.get(key)
        if root:
            candidates.append(Path(root) / "LibreOffice" / "program" / "soffice.exe")
    candidates.extend(
        [
            Path("/usr/bin/soffice"),
            Path("/usr/local/bin/soffice"),
            Path("/opt/libreoffice/program/soffice"),
            Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
        ]
    )
    return candidates


def find_soffice() -> Path:
    for candidate in soffice_candidates():
        try:
            if candidate.is_file():
                return candidate.resolve()
        except OSError:
            continue
    raise PresentationError(
        "LibreOffice soffice was not found. Set ARCFORGE_SOFFICE_PATH or install "
        "LibreOffice only after user approval."
    )


def render_pdf(input_path: Path, output_path: Path, force: bool) -> Dict[str, Any]:
    if not input_path.is_file():
        raise PresentationError("Input presentation does not exist: " + str(input_path))
    if output_path.exists() and not force:
        raise PresentationError(
            "Output already exists. Use a new path or pass --force only after explicit approval: "
            + str(output_path)
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    soffice = find_soffice()
    with tempfile.TemporaryDirectory(prefix="arcforge-slides-") as temporary_dir:
        temporary_root = Path(temporary_dir)
        profile = temporary_root / "profile"
        profile.mkdir()
        command = [
            str(soffice),
            "--headless",
            "-env:UserInstallation=" + profile.resolve().as_uri(),
            "--convert-to",
            "pdf",
            "--outdir",
            str(temporary_root),
            str(input_path),
        ]
        try:
            completed = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
                timeout=180,
            )
        except (OSError, subprocess.TimeoutExpired) as error:
            raise PresentationError(
                "Failed to run LibreOffice: " + str(error)
            ) from error
        if completed.returncode != 0:
            detail = (completed.stderr or completed.stdout).strip()
            raise PresentationError("LibreOffice PDF conversion failed: " + detail)
        generated = temporary_root / (input_path.stem + ".pdf")
        if not generated.is_file():
            raise PresentationError("LibreOffice did not produce the expected PDF")
        os.replace(generated, output_path)
    return {
        "path": str(output_path),
        "size_bytes": output_path.stat().st_size,
        "renderer": str(soffice),
        "visually_rendered": True,
    }


def apply_template_override(spec: Dict[str, Any], template: Optional[str]) -> Dict[str, Any]:
    if not template:
        return spec
    if not is_svg_deck_spec(spec):
        raise PresentationError("--template is only supported for schema_version 3 SVG deck manifests")
    updated = dict(spec)
    updated["mode"] = "template"
    updated["template"] = str(Path(template).expanduser().resolve())
    return updated


def resolve_asset_cache(args: argparse.Namespace) -> Optional[Path]:
    raw = getattr(args, "asset_cache", None)
    if not raw:
        return None
    cache = Path(str(raw)).expanduser().resolve()
    if not cache.is_dir():
        raise PresentationError("--asset-cache must be an existing directory: " + str(cache))
    return cache


def run_create(args: argparse.Namespace) -> Dict[str, Any]:
    spec, spec_path = load_json_object(args.spec)
    output_path = normalized_output_path(args.output, ".pptx")
    spec = apply_template_override(spec, getattr(args, "template", None))
    if is_svg_deck_spec(spec):
        presentation, deck = create_svg_deck(spec, spec_path.parent, resolve_asset_cache(args))
        atomic_save(presentation, output_path, args.force)
        write_build_stamp(spec_path.parent, deck.pop("slide_fingerprints", {}), output_path)
        return {
            "action": "created",
            "deck": deck,
            "presentation": inspect_presentation(output_path),
        }
    presentation = create_presentation(spec, spec_path.parent)
    atomic_save(presentation, output_path, args.force)
    return {"action": "created", "presentation": inspect_presentation(output_path)}


def run_validate(args: argparse.Namespace) -> Dict[str, Any]:
    """Convert an SVG deck manifest in memory and report every issue without writing a file."""
    spec, spec_path = load_json_object(args.spec)
    spec = apply_template_override(spec, getattr(args, "template", None))
    if not is_svg_deck_spec(spec):
        raise PresentationError("validate requires a schema_version 3 SVG deck manifest")
    _presentation, deck = create_svg_deck(spec, spec_path.parent, resolve_asset_cache(args))
    deck.pop("slide_fingerprints", None)
    issues = (
        len(deck["text_overflows"])
        + len(deck["out_of_bounds"])
        + len(deck["protected_collisions"])
    )
    return {
        "action": "validated",
        "valid": True,
        "issue_count": issues,
        "deck": deck,
    }


def run_inspect(args: argparse.Namespace) -> Dict[str, Any]:
    input_path = Path(args.input).expanduser().resolve()
    return {"action": "inspected", "presentation": inspect_presentation(input_path)}


def run_render(args: argparse.Namespace) -> Dict[str, Any]:
    input_path = Path(args.input).expanduser().resolve()
    output_path = normalized_output_path(args.output, ".pdf")
    return {
        "action": "rendered",
        "input": str(input_path),
        "pdf": render_pdf(input_path, output_path, args.force),
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Create, inspect, and optionally render PowerPoint decks for ArcForge."
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    create_parser = subparsers.add_parser("create", help="Create a PPTX deck from JSON")
    create_parser.add_argument("--spec", required=True, help="UTF-8 JSON specification")
    create_parser.add_argument("--output", required=True, help="Destination .pptx path")
    create_parser.add_argument(
        "--template",
        help="Existing .pptx used as the base deck (masters, layouts, theme) for SVG deck manifests",
    )
    create_parser.add_argument(
        "--asset-cache",
        help="Directory with SVG assets normalized by ArcForge (.arcforge-assets next to the manifest)",
    )
    create_parser.add_argument(
        "--force", action="store_true", help="Overwrite the exact output path"
    )
    create_parser.set_defaults(handler=run_create)

    validate_parser = subparsers.add_parser(
        "validate", help="Convert an SVG deck manifest in memory and report layout issues"
    )
    validate_parser.add_argument("--spec", required=True, help="UTF-8 JSON deck manifest")
    validate_parser.add_argument("--template", help="Existing .pptx used as the base deck")
    validate_parser.add_argument(
        "--asset-cache",
        help="Directory with SVG assets normalized by ArcForge (.arcforge-assets next to the manifest)",
    )
    validate_parser.set_defaults(handler=run_validate)

    inspect_parser = subparsers.add_parser(
        "inspect", help="Print a structural deck summary"
    )
    inspect_parser.add_argument("--input", required=True, help="Existing .pptx path")
    inspect_parser.set_defaults(handler=run_inspect)

    render_parser = subparsers.add_parser(
        "render", help="Render a PPTX deck to PDF with LibreOffice"
    )
    render_parser.add_argument("--input", required=True, help="Existing .pptx path")
    render_parser.add_argument("--output", required=True, help="Destination .pdf path")
    render_parser.add_argument(
        "--force", action="store_true", help="Overwrite the exact output path"
    )
    render_parser.set_defaults(handler=run_render)
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        require_pptx()
        result = args.handler(args)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except PresentationError as error:
        print("error: " + str(error), file=sys.stderr)
        return 2
    except Exception as error:
        print("error: unexpected presentation failure: " + str(error), file=sys.stderr)
        return 3


if __name__ == "__main__":
    raise SystemExit(main())
