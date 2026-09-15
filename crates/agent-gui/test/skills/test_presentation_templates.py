"""Round-trip regression tests for uploaded-template reuse and preview inheritance."""
import importlib.util
import io
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

SCRIPT = Path(__file__).resolve().parents[2] / "src-tauri/prompt/skills/arcforge-slides/scripts/presentation.py"
SPEC = importlib.util.spec_from_file_location("template_presentation", SCRIPT)
p = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = p
SPEC.loader.exec_module(p)


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.template = self.root / "template.pptx"
        deck = p.Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        text = slide.shapes.add_textbox(p.Inches(1), p.Inches(1), p.Inches(5), p.Inches(1))
        text.text = "Original title"
        text.text_frame.paragraphs[0].runs[0].font.size = p.Pt(30)
        self.title_id = text.shape_id
        table = slide.shapes.add_table(2, 2, p.Inches(1), p.Inches(3), p.Inches(5), p.Inches(1))
        self.table_id = table.shape_id
        table.table.cell(0, 0).text = "Original cell"
        slide.notes_slide.notes_text_frame.text = "Original speaker notes"
        data = p.ChartData()
        data.categories = ["A", "B"]
        data.add_series("Revenue", [1, 2])
        slide.shapes.add_chart(p.XL_CHART_TYPE.COLUMN_CLUSTERED, p.Inches(7), p.Inches(1), p.Inches(2), p.Inches(2), data)
        png = self.root / "image.png"
        p.Image.new("RGB", (12, 12), "blue").save(png)
        group = slide.shapes.add_group_shape()
        picture = group.shapes.add_picture(str(png), p.Inches(1), p.Inches(5))
        from pptx.opc.package import Part
        svg = Part(p.PackURI("/ppt/media/vector.svg"), "image/svg+xml", deck.part.package,
                   b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 12 12"><rect width="12" height="12" fill="blue"/></svg>')
        svg_id = slide.part.relate_to(svg, p.REL_NS + "/image")
        blip = picture._element.find(".//" + p.qn("a:blip"))
        extension = p.etree.SubElement(p.etree.SubElement(blip, p.qn("a:extLst")), p.qn("a:ext"), uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        p.etree.SubElement(extension, p.SVG_BLIP).set("{" + p.REL_NS + "}embed", svg_id)
        deck.save(self.template)

    def create(self, slides):
        return p.create_svg_deck({"schema_version": 3, "mode": "template", "template": "template.pptx", "slides": slides}, self.root)

    def test_duplicate_source_keeps_images_notes_charts_and_formatting(self):
        deck, report = self.create([
            {"slide_id": "one", "source_slide": 1,
             "text_edits": [{"shape_id": self.title_id, "text": "New title"}],
             "table_edits": [{"shape_id": self.table_id, "rows": [["Key", "Value"], ["A", "10"]]}]},
            {"slide_id": "two", "source_slide": 1},
        ])
        target = self.root / "output.pptx"
        p.atomic_save(deck, target, False)
        restored = p.Presentation(target)
        self.assertEqual(len(restored.slides), 2)
        self.assertEqual(p.relationship_errors(target), [])
        self.assertEqual(restored.slides[0].shapes[0].text, "New title")
        self.assertEqual(restored.slides[1].shapes[0].text, "Original title")
        self.assertEqual(restored.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt, 30)
        self.assertEqual(restored.slides[0].shapes[1].table.cell(1, 1).text, "10")
        for slide in restored.slides:
            self.assertEqual(slide.notes_slide.notes_text_frame.text, "Original speaker notes")
            note_slide = next(rel.target_part for rel in slide.notes_slide.part.rels.values() if rel.reltype.endswith("/slide"))
            self.assertIs(note_slide, slide.part)
        self.assertIsNot(restored.slides[0].shapes[2].chart.part, restored.slides[1].shapes[2].chart.part)
        with zipfile.ZipFile(target) as z:
            self.assertEqual(len(z.namelist()), len(set(z.namelist())))
            self.assertEqual(sum(name.startswith("ppt/slides/") and name.endswith(".xml") and "/_rels/" not in name for name in z.namelist()), 2)
        summary = p.inspect_presentation(target)
        self.assertEqual(summary["relationship_errors"], [])
        self.assertEqual(summary["slides"][0]["elements"][0]["shape_id"], self.title_id)
        self.assertEqual(report["template_slides_removed"], 1)

    def test_relationship_swap_is_simultaneous_and_bad_svg_is_rejected(self):
        root = p.etree.fromstring(f'<x xmlns:r="{p.REL_NS}"><png r:embed="rId5"/><svg r:embed="rId6"/></x>')
        p.remap_xml_relationships(root, {"rId6": "rId5", "rId5": "rId6"})
        self.assertEqual([node.get("{" + p.REL_NS + "}embed") for node in root], ["rId6", "rId5"])
        deck = p.Presentation(self.template)
        picture = list(deck.slides[0].shapes)[-1].shapes[0]
        blip = picture._element.find(".//" + p.qn("a:blip"))
        blip.find(".//" + p.SVG_BLIP).set("{" + p.REL_NS + "}embed", blip.get("{" + p.REL_NS + "}embed"))
        deck.save(self.template)
        self.assertEqual(len(p.relationship_errors(self.template)), 1)
        with self.assertRaisesRegex(p.PresentationError, "SVG reference"):
            self.create([{"slide_id": "bad", "source_slide": 1}])

    def test_invalid_targets_and_table_sizes_fail(self):
        for edit in ([{"shape_id": 999, "text": "x"}], [{"shape_id": self.title_id, "shape_path": [999, self.title_id], "text": "x"}]):
            with self.assertRaisesRegex(p.PresentationError, "exactly one"):
                self.create([{"slide_id": "bad", "source_slide": 1, "text_edits": edit}])
        with self.assertRaisesRegex(p.PresentationError, "dimensions"):
            self.create([{"slide_id": "bad", "source_slide": 1, "table_edits": [{"shape_id": self.table_id, "rows": [["x"]]}]}])
        with self.assertRaisesRegex(p.PresentationError, "page count"):
            self.create([{"slide_id": "bad", "source_slide": 2}])

    def test_missing_relationship_file_cannot_hide_dangling_chart_data(self):
        with zipfile.ZipFile(self.template) as source:
            payload = {name: source.read(name) for name in source.namelist()
                       if name != "ppt/charts/_rels/chart1.xml.rels"}
        with zipfile.ZipFile(self.template, "w", zipfile.ZIP_DEFLATED) as target:
            for name, data in payload.items():
                target.writestr(name, data)
        errors = p.relationship_errors(self.template)
        self.assertTrue(any(error["part"] == "ppt/charts/chart1.xml" and
                            error["error"] == "undefined relationship id" for error in errors))
        with self.assertRaisesRegex(p.PresentationError, "undefined relationship id"):
            self.create([{"slide_id": "bad", "source_slide": 1}])

    def test_per_page_layout_and_edit_fingerprints(self):
        source = [{"slide_id": "one", "source_slide": 1, "layout": 0}]
        deck, first = self.create(source)
        self.assertEqual(deck.slides[0].slide_layout.name, deck.slide_layouts[0].name)
        source[0]["text_edits"] = [{"shape_id": self.title_id, "text": "Changed"}]
        _, second = self.create(source)
        self.assertNotEqual(first["slide_fingerprints"], second["slide_fingerprints"])

    def test_custom_show_cannot_keep_or_play_original_template_pages(self):
        deck = p.Presentation(self.template)
        old_id = deck.slides._sldIdLst[0]
        shows = p.etree.SubElement(deck._element, p.qn("p:custShowLst"))
        show = p.etree.SubElement(shows, p.qn("p:custShow"), name="Original content", id="0")
        slides = p.etree.SubElement(show, p.qn("p:sldLst"))
        p.etree.SubElement(slides, p.qn("p:sld")).set("{" + p.REL_NS + "}id", old_id.rId)
        properties = p.etree.SubElement(deck._element, p.qn("p:showPr"))
        p.etree.SubElement(properties, p.qn("p:custShow"), id="0")
        extension = p.etree.SubElement(p.etree.SubElement(deck._element, p.qn("p:extLst")), p.qn("p:ext"), uri="sections")
        section_ns = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
        sections = p.etree.SubElement(extension, section_ns + "sectionLst")
        p.etree.SubElement(p.etree.SubElement(sections, section_ns + "section", name="Old section"), section_ns + "sldIdLst")
        deck.save(self.template)
        result, _ = self.create([{"slide_id": "new", "source_slide": 1,
                                  "text_edits": [{"shape_id": self.title_id, "text": "Updated title"}]}])
        target = self.root / "custom-show-result.pptx"
        p.atomic_save(result, target, False)
        with zipfile.ZipFile(target) as archive:
            slide_parts = [name for name in archive.namelist() if name.startswith("ppt/slides/") and name.endswith(".xml") and "/_rels/" not in name]
            self.assertEqual(len(slide_parts), 1)
            self.assertNotIn(b"Original title", archive.read(slide_parts[0]))
            xml = p.etree.fromstring(archive.read("ppt/presentation.xml"))
            self.assertIsNone(xml.find(p.qn("p:custShowLst")))
            self.assertIsNotNone(xml.find(p.qn("p:showPr")).find(p.qn("p:sldAll")))
            self.assertIsNone(xml.find(".//" + section_ns + "sectionLst"))

    def test_preview_materializes_artwork_without_modifying_template(self):
        deck = p.Presentation(self.template)
        layout = deck.slides[0].slide_layout
        # Put an existing native shape on the layout, as corporate templates do.
        artwork = p.deepcopy(deck.slides[0].shapes[0]._element)
        layout.shapes._spTree.append(artwork)
        deck.save(self.template)
        original = self.template.read_bytes()
        target = self.root / "preview.pptx"
        result = p.prepare_preview_presentation(self.template, target)
        preview = p.Presentation(target)
        self.assertGreater(result["materialized_shapes"], 0)
        self.assertEqual(self.template.read_bytes(), original)
        self.assertFalse(any(not shape.is_placeholder for shape in preview.slides[0].slide_layout.shapes))
        self.assertEqual(preview.slides[0].shapes[0].text, "Original title")
        self.assertEqual(p.relationship_errors(target), [])


if __name__ == "__main__":
    unittest.main()
